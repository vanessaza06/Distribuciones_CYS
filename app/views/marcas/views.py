"""
Vistas de Marcas (reemplaza tu archivo completo).
Requiere: pip install openpyxl reportlab
"""
import base64
import io
import os

from django.conf import settings
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.html import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import Marca  # type: ignore


# ---------------------------------------------------------------------------
# Filtro compartido: lo usan el listado y las 3 exportaciones,
# así lo que ves en pantalla es exactamente lo que se exporta.
# ---------------------------------------------------------------------------
CANDIDATOS_FECHA = ('fecha_creacion', 'fecha_registro', 'created_at', 'fecha')


def _campo_fecha():
    """Devuelve (nombre_campo, es_datetime) o (None, False) si el modelo no tiene fecha."""
    for nombre in CANDIDATOS_FECHA:
        try:
            campo = Marca._meta.get_field(nombre)
        except Exception:
            continue
        return nombre, campo.get_internal_type() == 'DateTimeField'
    return None, False


def _filtrar_marcas(request):
    query = request.GET.get('q', '').strip()
    estado_filtro = request.GET.get('estado', 'todos')
    desde = request.GET.get('desde', '').strip()
    hasta = request.GET.get('hasta', '').strip()

    marcas = Marca.objects.annotate(
        total_productos=Count('productos')
    ).order_by('nombre')

    if query:
        marcas = marcas.filter(
            Q(nombre__icontains=query) | Q(descripcion__icontains=query)
        )

    if estado_filtro in ('activo', 'inactivo'):
        marcas = marcas.filter(estado=estado_filtro)

    f_desde = parse_date(desde) if desde else None
    f_hasta = parse_date(hasta) if hasta else None

    campo, es_dt = _campo_fecha()
    if campo:
        sufijo = '__date' if es_dt else ''
        if f_desde:
            marcas = marcas.filter(**{f'{campo}{sufijo}__gte': f_desde})
        if f_hasta:
            marcas = marcas.filter(**{f'{campo}{sufijo}__lte': f_hasta})

    filtros = {
        'q': query,
        'estado': estado_filtro,
        'desde': desde,
        'hasta': hasta,
    }
    return marcas, filtros


# ---------------------------------------------------------------------------
# LISTADO (también atiende ?export=excel|pdf|imprimir)
# ---------------------------------------------------------------------------
def lista_marcas(request):
    # ---- Exportación: el JS llama a esta misma URL con ?export=... ----
    tipo = request.GET.get('export')
    if tipo == 'excel':
        return exportar_marcas_excel(request)
    if tipo == 'pdf':
        return exportar_marcas_pdf(request)
    if tipo == 'imprimir':
        return exportar_marcas_imprimir(request)

    # ---- Listado normal ----
    marcas, filtros = _filtrar_marcas(request)

    # Una sola consulta para las 3 tarjetas (antes eran 3)
    stats = Marca.objects.aggregate(
        total=Count('pk'),
        activas=Count('pk', filter=Q(estado='activo')),
        inactivas=Count('pk', filter=Q(estado='inactivo')),
    )

    paginator = Paginator(marcas, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'marcas/marcas.html', {
        'marcas': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'total_marcas': stats['total'],
        'total_activas': stats['activas'],
        'total_inactivas': stats['inactivas'],
        'query': filtros['q'],
        'estado_filtro': filtros['estado'],
        'desde': filtros['desde'],
        'hasta': filtros['hasta'],
        'filtro_fechas': _campo_fecha()[0] is not None,  # habilita Desde / Hasta solo si hay campo de fecha
    })


# ---------------------------------------------------------------------------
# CRUD
# ---------------------------------------------------------------------------
def crear_marca(request):
    if request.method == 'POST':
        nombre = (request.POST.get('nombre') or '').strip()
        descripcion = (request.POST.get('descripcion') or '').strip()
        estado = request.POST.get('estado', 'activo')
        if estado not in ('activo', 'inactivo'):
            estado = 'activo'

        if not nombre:
            messages.error(request, 'El nombre de la marca es obligatorio.')
        else:
            Marca.objects.create(nombre=nombre, descripcion=descripcion, estado=estado)
            messages.success(request, f'La marca "{nombre}" se creó correctamente.')
    return redirect('lista_marcas')


def editar_marca(request, codigo_marca):
    marca = get_object_or_404(Marca, codigo_marca=codigo_marca)

    if request.method == 'POST':
        marca.nombre = request.POST.get('nombre')
        marca.descripcion = request.POST.get('descripcion', '')
        marca.estado = request.POST.get('estado', marca.estado)
        marca.save()
        messages.success(request, f'La marca "{marca.nombre}" se actualizó correctamente.')

    return redirect('lista_marcas')


def cambiar_estado_marca(request, codigo_marca):
    marca = get_object_or_404(Marca, codigo_marca=codigo_marca)

    if request.method == 'POST':
        if marca.estado == 'activo':
            marca.estado = 'inactivo'
            messages.success(request, f'La marca "{marca.nombre}" se desactivó correctamente.')
        else:
            marca.estado = 'activo'
            messages.success(request, f'La marca "{marca.nombre}" se activó correctamente.')
        marca.save()

    return redirect('lista_marcas')


# ---------------------------------------------------------------------------
# Identidad visual (colores tomados del logo de CYS Ltda)
# ---------------------------------------------------------------------------
NAVY = "0B1B30"      # azul noche del logo
GOLD = "C48F3C"      # dorado del logo
CREAM = "F0E7DA"     # crema del texto "CYS"
ZEBRA = "FAF5EA"     # fila alterna suave
BORDE = "E3D3B5"     # bordes finos dorados claros
VERDE = "1E8449"
ROJO = "B03A2E"

LOGO_ARCHIVO = "LOGO_NUEVO.png"   # <- nombre de tu logo dentro de la carpeta static
_LOGO_CACHE = {}


def _ruta_logo():
    """Busca el logo en tus carpetas static. Devuelve la ruta o None."""
    if "ruta" in _LOGO_CACHE:
        return _LOGO_CACHE["ruta"]

    from django.contrib.staticfiles import finders

    ruta = None
    for candidato in (LOGO_ARCHIVO, f"img/{LOGO_ARCHIVO}", f"images/{LOGO_ARCHIVO}",
                      f"imagenes/{LOGO_ARCHIVO}", f"img/logo/{LOGO_ARCHIVO}"):
        try:
            ruta = finders.find(candidato)
        except Exception:
            ruta = None
        if ruta:
            break

    if not ruta:
        base = getattr(settings, "BASE_DIR", ".")
        carpetas = [str(d) for d in getattr(settings, "STATICFILES_DIRS", [])]
        carpetas += [os.path.join(str(base), "static"), os.path.join(str(base), "app", "static")]
        for carpeta in carpetas:
            for raiz, _dirs, archivos in os.walk(carpeta):
                if LOGO_ARCHIVO in archivos:
                    ruta = os.path.join(raiz, LOGO_ARCHIVO)
                    break
            if ruta:
                break

    _LOGO_CACHE["ruta"] = ruta
    return ruta


def _logo_png(px):
    """Logo redondo, redimensionado, con esquinas transparentes. Devuelve bytes PNG o None."""
    clave = ("png", px)
    if clave in _LOGO_CACHE:
        return _LOGO_CACHE[clave]

    datos = None
    ruta = _ruta_logo()
    if ruta:
        try:
            from PIL import Image, ImageDraw
            img = Image.open(ruta).convert("RGBA")
            mascara = Image.new("L", img.size, 0)
            ImageDraw.Draw(mascara).ellipse((0, 0, img.size[0] - 1, img.size[1] - 1), fill=255)
            img.putalpha(mascara)
            img = img.resize((px, px), Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            datos = buf.getvalue()
        except Exception:
            datos = None

    _LOGO_CACHE[clave] = datos
    return datos


def _encabezados():
    base = ["#", "Marca", "Descripción", "Productos", "Estado"]
    if _campo_fecha()[0]:
        base.append("Fecha creación")
    return base


def _filas(qs):
    campo = _campo_fecha()[0]
    filas = []
    for i, m in enumerate(qs, start=1):
        fila = [
            i,
            m.nombre,
            m.descripcion or "—",
            m.total_productos,
            "Activa" if m.estado == "activo" else "Inactiva",
        ]
        if campo:
            valor = getattr(m, campo, None)
            fila.append(valor.strftime("%d/%m/%Y") if valor else "—")
        filas.append(fila)
    return filas


def _texto_filtros(f):
    partes = []
    if f["q"]:
        partes.append(f"Búsqueda: {f['q']}")
    if f["estado"] in ("activo", "inactivo"):
        partes.append("Estado: " + ("Activa" if f["estado"] == "activo" else "Inactiva"))
    if f["desde"]:
        partes.append(f"Desde: {f['desde']}")
    if f["hasta"]:
        partes.append(f"Hasta: {f['hasta']}")
    return " | ".join(partes) if partes else "Sin filtros"


def _sello():
    return timezone.localtime().strftime('%Y%m%d_%H%M')


def _resumen(filas):
    activas = sum(1 for f in filas if f[4] == "Activa")
    return len(filas), activas, len(filas) - activas


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------
def exportar_marcas_excel(request):
    qs, filtros = _filtrar_marcas(request)
    filas = _filas(qs)
    total, activas, inactivas = _resumen(filas)
    encabezados = _encabezados()
    n = len(encabezados)
    ultima = get_column_letter(n)

    wb = Workbook()
    ws = wb.active
    ws.title = "Marcas"
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_view.showGridLines = False

    # Anchos de columna
    anchos = [12, 28, 48, 13, 13, 16][:n]
    for idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    relleno_navy = PatternFill("solid", fgColor=NAVY)
    relleno_oro = PatternFill("solid", fgColor=GOLD)

    # --- Banner superior (filas 1-4) ---
    for r in range(1, 5):
        ws.row_dimensions[r].height = 24
        for c in range(1, n + 1):
            ws.cell(row=r, column=c).fill = relleno_navy

    ws.merge_cells(f"B1:{ultima}2")
    ws["B1"] = "CYS LTDA · DISTRIBUCIONES"
    ws["B1"].font = Font(name="Georgia", bold=True, size=20, color=GOLD)
    ws["B1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)

    ws.merge_cells(f"B3:{ultima}4")
    ws["B3"] = "Reporte de marcas"
    ws["B3"].font = Font(name="Calibri", size=14, color=CREAM)
    ws["B3"].alignment = Alignment(vertical="top", horizontal="left", indent=1)

    logo = _logo_png(200)
    if logo:
        img = XLImage(io.BytesIO(logo))
        img.width, img.height = 92, 92
        ws.add_image(img, "A1")

    # Franja dorada
    ws.row_dimensions[5].height = 5
    for c in range(1, n + 1):
        ws.cell(row=5, column=c).fill = relleno_oro

    # Filtros y fecha de generación
    ws.merge_cells(f"A6:{ultima}6")
    ws["A6"] = f"{_texto_filtros(filtros)}   ·   Generado: {timezone.localtime():%d/%m/%Y %H:%M}"
    ws["A6"].font = Font(italic=True, size=10, color="6B6B6B")
    ws["A6"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 6

    # --- Encabezado de la tabla ---
    fila_enc = 8
    ws.row_dimensions[fila_enc].height = 26
    borde_oro = Border(bottom=Side(style="medium", color=GOLD))
    for c, texto in enumerate(encabezados, start=1):
        cell = ws.cell(row=fila_enc, column=c, value=texto)
        cell.font = Font(bold=True, color=GOLD, size=11)
        cell.fill = relleno_navy
        cell.border = borde_oro
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- Datos ---
    borde_fino = Border(
        left=Side(style="thin", color=BORDE), right=Side(style="thin", color=BORDE),
        top=Side(style="thin", color=BORDE), bottom=Side(style="thin", color=BORDE),
    )
    for k, fila in enumerate(filas):
        r = fila_enc + 1 + k
        ws.row_dimensions[r].height = 21
        for c, valor in enumerate(fila, start=1):
            cell = ws.cell(row=r, column=c, value=valor)
            cell.border = borde_fino
            cell.font = Font(size=11, color="1F2937")
            if k % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
            if c in (1, 4, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        # Marca en negrita y estado con color
        ws.cell(row=r, column=2).font = Font(bold=True, size=11, color=NAVY)
        est = ws.cell(row=r, column=5)
        est.font = Font(bold=True, size=11, color=VERDE if fila[4] == "Activa" else ROJO)

    # --- Totales ---
    r_tot = fila_enc + len(filas) + 2
    ws.merge_cells(f"A{r_tot}:{ultima}{r_tot}")
    ws[f"A{r_tot}"] = f"Total: {total} marcas   ·   Activas: {activas}   ·   Inactivas: {inactivas}"
    ws[f"A{r_tot}"].font = Font(bold=True, size=11, color=NAVY)
    ws[f"A{r_tot}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.row_dimensions[r_tot].height = 24
    for c in range(1, n + 1):
        ws.cell(row=r_tot, column=c).fill = PatternFill("solid", fgColor=CREAM)
        ws.cell(row=r_tot, column=c).border = Border(top=Side(style="medium", color=GOLD))

    # Extras: filtros, congelar encabezado, impresión
    if filas:
        ws.auto_filter.ref = f"A{fila_enc}:{ultima}{fila_enc + len(filas)}"
    ws.freeze_panes = f"A{fila_enc + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_enc}:{fila_enc}"
    ws.print_options.horizontalCentered = True

    buffer = io.BytesIO()
    wb.save(buffer)

    resp = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="marcas_{_sello()}.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def exportar_marcas_pdf(request):
    qs, filtros = _filtrar_marcas(request)
    filas = _filas(qs)
    total, activas, inactivas = _resumen(filas)
    encabezados = _encabezados()

    c_navy = colors.HexColor("#" + NAVY)
    c_gold = colors.HexColor("#" + GOLD)
    c_cream = colors.HexColor("#" + CREAM)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=40,
        title="Reporte de marcas — CYS Ltda",
    )
    ancho_util = landscape(A4)[0] - 60

    estilos = getSampleStyleSheet()
    st_titulo = ParagraphStyle("t", parent=estilos["Title"], fontName="Times-Bold", fontSize=22,
                               textColor=c_gold, alignment=0, spaceAfter=2, leading=24)
    st_sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=12, textColor=c_cream, leading=15)
    st_meta = ParagraphStyle("m", parent=estilos["Normal"], fontSize=8.5,
                             textColor=colors.HexColor("#6B6B6B"), leading=12)
    st_celda = ParagraphStyle("c", parent=estilos["BodyText"], fontSize=9, leading=11,
                              textColor=colors.HexColor("#1F2937"))
    st_marca = ParagraphStyle("mk", parent=st_celda, fontName="Helvetica-Bold", textColor=c_navy)
    st_centro = ParagraphStyle("ce", parent=st_celda, alignment=TA_CENTER)

    # --- Banner con logo ---
    logo = _logo_png(300)
    celda_logo = RLImage(io.BytesIO(logo), width=62, height=62) if logo else ""
    celda_texto = [Paragraph("CYS LTDA · DISTRIBUCIONES", st_titulo),
                   Paragraph("Reporte de marcas", st_sub)]
    banner = Table([[celda_logo, celda_texto]], colWidths=[80, ancho_util - 80])
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), c_navy),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, -1), 3, c_gold),
    ]))

    elementos = [
        banner,
        Spacer(1, 8),
        Paragraph(escape(_texto_filtros(filtros)) + " &nbsp;·&nbsp; Generado: "
                  + timezone.localtime().strftime("%d/%m/%Y %H:%M"), st_meta),
        Spacer(1, 10),
    ]

    # --- Tabla ---
    st_head = ParagraphStyle("h", parent=st_celda, fontName="Helvetica-Bold", textColor=c_gold,
                             alignment=TA_CENTER, fontSize=9.5)
    data = [[Paragraph(h, st_head) for h in encabezados]]
    for f in filas:
        color = "#" + (VERDE if f[4] == "Activa" else ROJO)
        fila = [
            Paragraph(str(f[0]), st_centro),
            Paragraph(escape(f[1]), st_marca),
            Paragraph(escape(f[2]), st_celda),
            Paragraph(str(f[3]), st_centro),
            Paragraph(f'<font color="{color}"><b>{f[4]}</b></font>', st_centro),
        ]
        if len(f) > 5:
            fila.append(Paragraph(str(f[5]), st_centro))
        data.append(fila)

    if not filas:
        data.append([Paragraph("No hay marcas para mostrar.", st_centro)] + [""] * (len(encabezados) - 1))

    anchos = [30, 150, 290, 65, 70, 90] if len(encabezados) == 6 else [30, 170, 340, 70, 80]
    escala = ancho_util / sum(anchos)
    tabla = Table(data, colWidths=[a * escala for a in anchos], repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), c_navy),
        ("LINEBELOW", (0, 0), (-1, 0), 2, c_gold),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 1), (-1, -1), 0.4, colors.HexColor("#" + BORDE)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + ZEBRA)]),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]
    if not filas:
        estilo.append(("SPAN", (0, 1), (-1, 1)))
    tabla.setStyle(TableStyle(estilo))
    elementos.append(tabla)

    # --- Totales ---
    st_tot = ParagraphStyle("tot", parent=st_celda, alignment=2, fontName="Helvetica-Bold",
                            textColor=c_navy, fontSize=10)
    resumen = Table(
        [[Paragraph(f"Total: {total} marcas &nbsp;·&nbsp; Activas: {activas} &nbsp;·&nbsp; Inactivas: {inactivas}", st_tot)]],
        colWidths=[ancho_util],
    )
    resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), c_cream),
        ("LINEABOVE", (0, 0), (-1, 0), 2, c_gold),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
    ]))
    elementos += [Spacer(1, 6), resumen]

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setStrokeColor(c_gold)
        canvas.setLineWidth(1)
        canvas.line(30, 28, landscape(A4)[0] - 30, 28)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6B6B6B"))
        canvas.drawString(30, 16, "CYS Ltda · Distribuciones")
        canvas.drawRightString(landscape(A4)[0] - 30, 16, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(elementos, onFirstPage=pie, onLaterPages=pie)

    resp = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="marcas_{_sello()}.pdf"'
    return resp


# ---------------------------------------------------------------------------
# IMPRIMIR
# ---------------------------------------------------------------------------
def exportar_marcas_imprimir(request):
    qs, filtros = _filtrar_marcas(request)
    filas = _filas(qs)
    total, activas, inactivas = _resumen(filas)
    encab = _encabezados()

    def fila_html(f):
        clase = "ok" if f[4] == "Activa" else "no"
        extra = "".join(f"<td class='c'>{escape(str(x))}</td>" for x in f[5:])
        return (
            "<tr>"
            f"<td class='c'>{f[0]}</td><td class='marca'>{escape(f[1])}</td><td>{escape(f[2])}</td>"
            f"<td class='c'>{f[3]}</td><td class='c {clase}'>{f[4]}</td>{extra}"
            "</tr>"
        )

    filas_html = "".join(fila_html(f) for f in filas) or \
        f"<tr><td colspan='{len(encab)}' class='c'>No hay marcas para mostrar.</td></tr>"
    encabezados = "".join(f"<th>{h}</th>" for h in encab)

    logo = _logo_png(300)
    logo_html = ""
    if logo:
        logo_html = f"<img class='logo' src='data:image/png;base64,{base64.b64encode(logo).decode()}' alt='CYS Ltda'>"

    html = f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><title>Reporte de marcas — CYS Ltda</title>
<style>
  @page {{ margin: 14mm; }}
  * {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
  body {{ font-family: Arial, sans-serif; margin: 0; color: #1f2937; }}
  .banner {{ background: #{NAVY}; border-bottom: 4px solid #{GOLD}; padding: 14px 20px;
            display: flex; align-items: center; gap: 18px; border-radius: 6px 6px 0 0; }}
  .logo {{ width: 70px; height: 70px; }}
  .banner h1 {{ margin: 0; font: bold 24px Georgia, 'Times New Roman', serif; color: #{GOLD}; letter-spacing: 1px; }}
  .banner p {{ margin: 3px 0 0; font-size: 14px; color: #{CREAM}; }}
  .meta {{ font-size: 11px; color: #6b6b6b; margin: 10px 2px 12px; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  th {{ background: #{NAVY}; color: #{GOLD}; padding: 9px 10px; border-bottom: 3px solid #{GOLD}; text-align: center; }}
  td {{ padding: 8px 10px; border: 1px solid #{BORDE}; }}
  td.c {{ text-align: center; }}
  td.marca {{ font-weight: bold; color: #{NAVY}; }}
  td.ok {{ color: #{VERDE}; font-weight: bold; }}
  td.no {{ color: #{ROJO}; font-weight: bold; }}
  tr:nth-child(even) td {{ background: #{ZEBRA}; }}
  .total {{ margin-top: 8px; background: #{CREAM}; border-top: 3px solid #{GOLD}; padding: 9px 14px;
           text-align: right; font-weight: bold; color: #{NAVY}; font-size: 13px; }}
  .pie {{ margin-top: 14px; font-size: 10px; color: #888; text-align: center; }}
</style></head>
<body>
  <div class="banner">
    {logo_html}
    <div><h1>CYS LTDA · DISTRIBUCIONES</h1><p>Reporte de marcas</p></div>
  </div>
  <div class="meta">{escape(_texto_filtros(filtros))} &nbsp;·&nbsp; Generado: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}</div>
  <table><thead><tr>{encabezados}</tr></thead><tbody>{filas_html}</tbody></table>
  <div class="total">Total: {total} marcas &nbsp;·&nbsp; Activas: {activas} &nbsp;·&nbsp; Inactivas: {inactivas}</div>
  <div class="pie">CYS Ltda · Distribuciones</div>
  <script>window.onload = function () {{ window.print(); }};</script>
</body></html>"""
    return HttpResponse(html)