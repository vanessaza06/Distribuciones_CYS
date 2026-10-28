from io import BytesIO
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.html import escape
from django.utils.http import url_has_allowed_host_and_scheme

from app.forms import PresentacionForm
from app.models import PresentacionProducto, Producto


def _volver(request, producto_pk, resaltar=None):
    """Si el formulario manda 'next' (ej: desde la página de Presentaciones),
    vuelve ahí; si no, mantiene el comportamiento anterior.
    Con 'resaltar' agrega ?resaltar=<pk> para marcar esa fila en la tabla."""
    destino = request.POST.get('next', '')
    if destino and url_has_allowed_host_and_scheme(
        destino,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        if resaltar:
            separador = '&' if '?' in destino else '?'
            destino = f"{destino}{separador}{urlencode({'resaltar': resaltar})}"
        return redirect(destino)
    return redirect('producto_detalle', pk=producto_pk)


@login_required
def presentacion_crear(request, producto_pk):
    producto = get_object_or_404(Producto, pk=producto_pk)

    if request.method != 'POST':
        return redirect('producto_detalle', pk=producto_pk)

    form = PresentacionForm(request.POST)

    if not form.is_valid():
        for campo, errores in form.errors.items():
            for error in errores:
                messages.error(request, f'{campo}: {error}')
        return _volver(request, producto_pk)

    nombre = form.cleaned_data['nombre'].strip()
    if producto.presentaciones.filter(nombre__iexact=nombre).exists():
        messages.error(request, f'Ya existe una presentación "{nombre}" en este producto.')
        return _volver(request, producto_pk)

    presentacion = form.save(commit=False)
    presentacion.producto = producto
    presentacion.save()

    messages.success(request, f'Presentación "{presentacion.nombre}" agregada.')
    return _volver(request, producto_pk, resaltar=presentacion.pk)


@login_required
def presentacion_editar(request, pk):
    pres = get_object_or_404(PresentacionProducto, pk=pk)

    if request.method != 'POST':
        return redirect('producto_detalle', pk=pres.producto_id)

    form = PresentacionForm(request.POST, instance=pres)

    if not form.is_valid():
        for campo, errores in form.errors.items():
            for error in errores:
                messages.error(request, f'{campo}: {error}')
        return _volver(request, pres.producto_id)

    nombre_nuevo = form.cleaned_data['nombre'].strip()
    duplicado = (
        pres.producto.presentaciones
        .filter(nombre__iexact=nombre_nuevo)
        .exclude(pk=pres.pk)
        .exists()
    )
    if duplicado:
        messages.error(request, f'Ya existe otra presentación "{nombre_nuevo}" en este producto.')
        return _volver(request, pres.producto_id)

    # OJO: ya NO se copia precio_venta al costo_unitario de los lotes.
    form.save()

    messages.success(request, f'Presentación "{pres.nombre}" actualizada.')
    return _volver(request, pres.producto_id, resaltar=pres.pk)


@login_required
def presentacion_toggle_activo(request, pk):
    # TODO: pasar a @require_POST cuando el template/JS envíe POST con CSRF.
    pres = get_object_or_404(PresentacionProducto, pk=pk)
    pres.activo = not pres.activo
    pres.save(update_fields=['activo'])
    estado = 'activada' if pres.activo else 'desactivada'
    messages.success(request, f'Presentación "{pres.nombre}" {estado}.')
    return _volver(request, pres.producto_id, resaltar=pres.pk)


@login_required
def presentacion_lista(request):
    presentaciones = (
        PresentacionProducto.objects
        .select_related('producto', 'producto__categoria')
        .order_by('producto__nombre', 'nombre')
    )
    productos = Producto.objects.filter(activo=True).order_by('nombre')

    resaltar = request.GET.get('resaltar', '')
    resaltar_pk = int(resaltar) if resaltar.isdigit() else None

    context = {
        'presentaciones': presentaciones,
        'productos': productos,
        'resaltar_pk': resaltar_pk,
        'breadcrumb_items': [
            {'nombre': 'Inventario', 'url': None},
            {'nombre': 'Presentaciones', 'url': None},
        ],
    }
    return render(request, 'presentaciones/presentaciones.html', context)


# ===============================
# EXPORTAR PRESENTACIONES (PDF / EXCEL / IMPRIMIR)
# ===============================
ENCABEZADOS_EXPORT_PRES = [
    'Producto', 'Categoría', 'Presentación', 'Cantidad',
    'Precio venta', 'Stock real', 'Estado',
]


def _filas_presentaciones():
    """Devuelve las filas (listas) de todas las presentaciones para exportar."""
    presentaciones = (
        PresentacionProducto.objects
        .select_related('producto', 'producto__categoria')
        .order_by('producto__nombre', 'nombre')
    )
    filas = []
    for p in presentaciones:
        filas.append([
            p.producto.nombre,
            p.producto.categoria.nombre if p.producto.categoria else '—',
            p.nombre,
            p.cantidad,
            p.precio_venta or 0,
            p.stock_real or 0,
            'Activa' if p.activo else 'Inactiva',
        ])
    return filas


def _nombre_archivo_pres(extension):
    return f'presentaciones_{timezone.localdate():%Y%m%d}.{extension}'


def _exportar_excel_presentaciones():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Presentaciones'

    ws.append(ENCABEZADOS_EXPORT_PRES)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0D2A4F')
        celda.alignment = Alignment(horizontal='center', vertical='center')

    for fila in _filas_presentaciones():
        fila[4] = float(fila[4] or 0)
        ws.append(fila)

    # Formato moneda y anchos de columna
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for celda in row:
            celda.number_format = '"$"#,##0'
    for i, titulo in enumerate(ENCABEZADOS_EXPORT_PRES, start=1):
        maximo = max(
            [len(str(titulo))] +
            [len(str(c.value or '')) for c in ws[get_column_letter(i)][1:]]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(maximo + 3, 40)
    ws.freeze_panes = 'A2'

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo_pres("xlsx")}"'
    return response


def _exportar_pdf_presentaciones():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm,
        title='Reporte de Presentaciones',
    )
    estilos = getSampleStyleSheet()

    datos = [ENCABEZADOS_EXPORT_PRES]
    for f in _filas_presentaciones():
        f = list(f)
        f[4] = f'${(f[4] or 0):,.0f}'
        datos.append(f)

    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D2A4F')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (3, 1), (5, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF3FA')]),
    ]))

    elementos = [
        Paragraph('Reporte de Presentaciones — CYS Ltda', estilos['Title']),
        Paragraph(f'Generado el {timezone.localdate():%d/%m/%Y}', estilos['Normal']),
        Spacer(1, 0.4 * cm),
        tabla,
    ]
    doc.build(elementos)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo_pres("pdf")}"'
    return response


def _exportar_imprimir_presentaciones():
    filas = _filas_presentaciones()

    encabezados_html = ''.join(f'<th>{escape(e)}</th>' for e in ENCABEZADOS_EXPORT_PRES)

    filas_html = ''
    for f in filas:
        filas_html += (
            '<tr>'
            f'<td>{escape(f[0])}</td>'
            f'<td>{escape(f[1])}</td>'
            f'<td>{escape(f[2])}</td>'
            f'<td class="num">{f[3]}</td>'
            f'<td class="num">${(f[4] or 0):,.0f}</td>'
            f'<td class="num">{f[5]}</td>'
            f'<td>{escape(f[6])}</td>'
            '</tr>'
        )
    if not filas_html:
        filas_html = '<tr><td colspan="7">No hay presentaciones registradas.</td></tr>'

    plural = '' if len(filas) == 1 else 'es'
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <title>Reporte de Presentaciones — CYS Ltda</title>
  <style>
    body {{ font-family: Arial, sans-serif; font-size: 12px; color: #111; margin: 24px; }}
    h1 {{ font-size: 20px; margin: 0 0 4px; }}
    .meta {{ color: #555; margin-bottom: 16px; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th {{ background: #0d2a4f; color: #fff; padding: 6px; text-align: left; }}
    td {{ border-bottom: 1px solid #ccc; padding: 5px 6px; }}
    tr:nth-child(even) td {{ background: #f3f6fb; }}
    td.num {{ text-align: right; }}
    @media print {{
      body {{ margin: 0; }}
      th {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
    }}
  </style>
</head>
<body>
  <h1>Reporte de Presentaciones — CYS Ltda</h1>
  <div class="meta">Generado el {timezone.localdate():%d/%m/%Y} · {len(filas)} presentaci{'ón' if len(filas) == 1 else 'ones'}</div>
  <table>
    <thead><tr>{encabezados_html}</tr></thead>
    <tbody>{filas_html}</tbody>
  </table>
  <script>window.onload = function () {{ window.print(); }};</script>
</body>
</html>"""
    return HttpResponse(html)


@login_required
def presentacion_exportar(request, formato):
    if formato == 'excel':
        return _exportar_excel_presentaciones()
    if formato == 'pdf':
        return _exportar_pdf_presentaciones()
    if formato == 'imprimir':
        return _exportar_imprimir_presentaciones()
    raise Http404('Formato de exportación no válido.')