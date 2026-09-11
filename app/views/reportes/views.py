"""
Vistas del módulo de Reportes y Exportación (PDF / Excel).
Distribuciones CYS
"""

import json
import zoneinfo
from io import BytesIO
from decimal import Decimal

from django.shortcuts import render
from django.utils import timezone
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.core.serializers.json import DjangoJSONEncoder
from django.urls import reverse
from django.db.models import Sum, Count, Q

# Modelos del sistema
from app.models import Venta, DetalleVenta, Devolucion, Compra, DetalleCompra, Proveedor, Producto, Lote

# ReportLab (PDF)
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

# openpyxl (Excel)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ZONA_COLOMBIA = zoneinfo.ZoneInfo('America/Bogota')

# ─────────────────────────────────────────────
#  COLORES CORPORATIVOS CYS
# ─────────────────────────────────────────────
COLOR_PRIMARIO   = colors.HexColor('#1A2B3C')   # azul oscuro
COLOR_ACENTO     = colors.HexColor('#4DA8DA')   # azul claro
COLOR_VERDE      = colors.HexColor('#2ecc71')
COLOR_AMARILLO   = colors.HexColor('#f1c40f')
COLOR_ROJO       = colors.HexColor('#e74c3c')
COLOR_GRIS_FILA  = colors.HexColor('#F0F4F8')
COLOR_BLANCO     = colors.white
COLOR_TEXTO      = colors.HexColor('#1A2B3C')

# Colores Excel
XLS_ROJO_MARCA   = 'A3242D'
XLS_ROJO_HEADER  = 'C0392B'
XLS_ROJO_SUAVE   = 'FBE4E4'
XLS_ROJO_BORDE   = 'E8B4B4'
XLS_ROJO_KPI_BG  = 'FDF1F1'
XLS_GRIS_TEXTO   = '4A4A4A'
XLS_BLANCO       = 'FFFFFF'
XLS_VERDE_OK     = '2E8B57'
XLS_ROJO_OUT     = 'C0392B'
XLS_FONT         = 'Calibri'

TIPOS_VALIDOS = {'ventas', 'proveedores', 'resumen_diario', 'analisis', 'devoluciones', 'compras'}


# ─────────────────────────────────────────────
#  HELPERS PDF
# ─────────────────────────────────────────────
def _build_pdf(titulo, subtitulo, elementos, orientacion='portrait'):
    buffer = BytesIO()
    pagesize = landscape(letter) if orientacion == 'landscape' else letter
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=1.5*cm,   bottomMargin=1.5*cm,
    )

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'Titulo', parent=styles['Title'],
        fontSize=18, textColor=COLOR_PRIMARIO,
        spaceAfter=2, alignment=TA_LEFT,
    )
    estilo_subtitulo = ParagraphStyle(
        'Sub', parent=styles['Normal'],
        fontSize=9, textColor=colors.HexColor('#6B8CA0'),
        spaceAfter=8, alignment=TA_LEFT,
    )

    historia = [
        Paragraph("CYS Ltda.", estilo_titulo),
        Paragraph(titulo, ParagraphStyle('T2', parent=styles['Heading1'], fontSize=13, textColor=COLOR_ACENTO, spaceAfter=2)),
        Paragraph(subtitulo, estilo_subtitulo),
        HRFlowable(width="100%", thickness=1.5, color=COLOR_ACENTO, spaceAfter=10)
    ]
    historia.extend(elementos)

    doc.build(historia)
    buffer.seek(0)
    return buffer


def _estilo_tabla_base(num_cols):
    return TableStyle([
        ('BACKGROUND',    (0, 0), (-1, 0), COLOR_PRIMARIO),
        ('TEXTCOLOR',     (0, 0), (-1, 0), COLOR_BLANCO),
        ('FONTNAME',      (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',      (0, 0), (-1, 0), 9),
        ('ALIGN',         (0, 0), (-1, 0), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 7),
        ('TOPPADDING',    (0, 0), (-1, 0), 7),
        ('FONTNAME',      (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE',      (0, 1), (-1, -1), 8),
        ('ALIGN',         (0, 1), (-1, -1), 'LEFT'),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING',    (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('ROWBACKGROUNDS',(0, 1), (-1, -1), [COLOR_BLANCO, COLOR_GRIS_FILA]),
        ('GRID',          (0, 0), (-1, -1), 0.4, colors.HexColor('#D0DCE8')),
        ('BOX',           (0, 0), (-1, -1), 1,   COLOR_ACENTO),
    ])


def _kpi_table(items):
    filas_label = []
    filas_val   = []
    for label, valor, color_hex in items:
        filas_label.append(Paragraph(
            f'<font color="{color_hex}"><b>{valor}</b></font>',
            ParagraphStyle('kv', fontSize=14, alignment=TA_CENTER)
        ))
        filas_val.append(Paragraph(
            label,
            ParagraphStyle('kl', fontSize=7, textColor=colors.HexColor('#6B8CA0'), alignment=TA_CENTER)
        ))

    col_w = [4*cm] * len(items)
    t = Table([filas_label, filas_val], colWidths=col_w)
    t.setStyle(TableStyle([
        ('BOX',           (0, 0), (-1, -1), 1,   colors.HexColor('#D0DCE8')),
        ('INNERGRID',     (0, 0), (-1, -1), 0.4, colors.HexColor('#D0DCE8')),
        ('BACKGROUND',    (0, 0), (-1, -1), COLOR_GRIS_FILA),
        ('TOPPADDING',    (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    return t


# ── GENERADORES PDF ──────────────────────────────────────────────────────────
def _pdf_ventas(ventas_qs, fecha_inicio, fecha_fin):
    elementos = []
    total_v = sum(v.total_venta for v in ventas_qs)
    total_u = sum(det.cantidad for v in ventas_qs for det in v.detalleventa_set.all())
    total_ventas_count = ventas_qs.count()

    elementos.append(_kpi_table([
        ('Total Ventas', f'${total_v:,.0f}', '#2ecc71'),
        ('Unidades',     str(total_u),        '#4DA8DA'),
        ('Órdenes',      str(total_ventas_count), '#f1c40f'),
    ]))
    elementos.append(Spacer(1, 12))

    cabecera = [['#', 'Fecha', 'Venta ID', 'Producto', 'Cant.', 'Precio Unit.', 'Total']]
    filas = cabecera
    contador = 1
    for v in ventas_qs:
        for det in v.detalleventa_set.all():
            filas.append([
                str(contador),
                v.fecha.astimezone(ZONA_COLOMBIA).strftime("%Y-%m-%d %H:%M"),
                f"#{v.codigo_venta}",
                det.producto.nombre if det.producto else '—',
                str(det.cantidad),
                f'${float(det.precio_unitario):,.0f}',
                f'${float(det.subtotal):,.0f}',
            ])
            contador += 1

    if len(filas) == 1:
        filas.append(['', 'Sin ventas en el período seleccionado', '', '', '', '', ''])

    col_w = [1*cm, 3.5*cm, 2.5*cm, 6*cm, 1.5*cm, 2.5*cm, 3*cm]
    t = Table(filas, colWidths=col_w, repeatRows=1)
    estilo = _estilo_tabla_base(7)
    estilo.add('ALIGN', (4, 1), (6, -1), 'RIGHT')
    t.setStyle(estilo)
    elementos.append(t)

    subtitulo = f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}  |  Generado: {timezone.now().astimezone(ZONA_COLOMBIA).strftime('%d/%m/%Y %H:%M')}"
    return _build_pdf("Historial de Ventas", subtitulo, elementos, orientacion='landscape')


def _pdf_proveedores(proveedores_list):
    elementos = []
    elementos.append(_kpi_table([
        ('Total Proveedores', str(len(proveedores_list)), '#5DCAA5'),
        ('Estado',            'Registrados',              '#4DA8DA'),
    ]))
    elementos.append(Spacer(1, 12))

    filas = [['#', 'NIT', 'Empresa', 'Teléfono', 'Correo', 'Estado']]
    for i, p in enumerate(proveedores_list, 1):
        filas.append([
            str(i),
            p.nit_proveedores,
            p.nombre_empresa,
            p.telefono or '—',
            p.correo,
            p.get_estado_display(),
        ])

    col_w = [1*cm, 3*cm, 6*cm, 3*cm, 5*cm, 2*cm]
    t = Table(filas, colWidths=col_w, repeatRows=1)
    t.setStyle(_estilo_tabla_base(6))
    elementos.append(t)

    subtitulo = f"Generado: {timezone.now().astimezone(ZONA_COLOMBIA).strftime('%d/%m/%Y %H:%M')}"
    return _build_pdf("Reporte de Proveedores", subtitulo, elementos)


def _pdf_resumen_diario(hoy, ventas_hoy, ingresos_hoy, top_productos_hoy):
    elementos = []
    total_uds_hoy = sum(det.cantidad for v in ventas_hoy for det in v.detalleventa_set.all())

    elementos.append(_kpi_table([
        ('Ingresos hoy',  f'${ingresos_hoy:,.0f}',    '#2ecc71'),
        ('Ventas hoy',    str(len(ventas_hoy)),       '#4DA8DA'),
        ('Und. Vendidas', str(total_uds_hoy),         '#f1c40f'),
    ]))
    elementos.append(Spacer(1, 12))

    filas_v = [['#', 'Venta', 'Producto', 'Cant.', 'Total', 'Hora']]
    cont = 1
    for v in ventas_hoy:
        for det in v.detalleventa_set.all():
            filas_v.append([
                str(cont), f"#{v.codigo_venta}", det.producto.nombre if det.producto else '—',
                str(det.cantidad), f'${float(det.subtotal):,.0f}',
                v.fecha.astimezone(ZONA_COLOMBIA).strftime("%H:%M"),
            ])
            cont += 1
    if len(filas_v) == 1:
        filas_v.append(['', 'Sin ventas hoy', '', '', '', ''])

    col_v = [1*cm, 2.5*cm, 7*cm, 2*cm, 3.5*cm, 2*cm]
    tv = Table(filas_v, colWidths=col_v, repeatRows=1)
    estilo_v = _estilo_tabla_base(6)
    estilo_v.add('ALIGN', (3, 1), (4, -1), 'RIGHT')
    tv.setStyle(estilo_v)
    elementos.append(tv)

    if top_productos_hoy:
        elementos.append(Spacer(1, 14))
        filas_t = [['#', 'Producto', 'Unidades vendidas', 'Total generado']]
        for i, (nombre, datos) in enumerate(top_productos_hoy, 1):
            filas_t.append([str(i), nombre, str(datos['cantidad']), f'${datos["subtotal"]:,.0f}'])
        col_t = [1*cm, 8*cm, 4*cm, 4*cm]
        tt = Table(filas_t, colWidths=col_t, repeatRows=1)
        tt.setStyle(_estilo_tabla_base(4))
        elementos.append(tt)

    subtitulo = f"Fecha: {hoy.strftime('%d/%m/%Y')}  |  Generado: {timezone.now().astimezone(ZONA_COLOMBIA).strftime('%H:%M')}"
    return _build_pdf("Resumen Diario", subtitulo, elementos)


def _pdf_devoluciones(devoluciones_qs, fecha_inicio, fecha_fin):
    elementos = []
    total_dev   = devoluciones_qs.count()
    total_valor = sum(d.total_devuelto for d in devoluciones_qs)

    elementos.append(_kpi_table([
        ('Devoluciones',   str(total_dev),          '#4DA8DA'),
        ('Total devuelto', f'${total_valor:,.0f}',  '#e74c3c'),
    ]))
    elementos.append(Spacer(1, 12))

    cabecera = [['#', 'Fecha', 'Venta Asoc.', 'Motivo', 'Tipo', 'Estado', 'Total devuelto']]
    filas = cabecera
    for i, d in enumerate(devoluciones_qs, 1):
        filas.append([
            str(i),
            d.fecha.astimezone(ZONA_COLOMBIA).strftime("%d/%m/%Y %H:%M"),
            f"Venta #{d.venta.codigo_venta}" if d.venta else '—',
            d.motivo[:25] + '...' if len(d.motivo) > 25 else d.motivo,
            d.tipo_devolucion,
            d.estado.capitalize(),
            f'${float(d.total_devuelto):,.0f}',
        ])

    if len(filas) == 1:
        filas.append(['', 'Sin devoluciones en el período seleccionado', '', '', '', '', ''])

    col_w = [1*cm, 3.5*cm, 3*cm, 4*cm, 3*cm, 2.5*cm, 3*cm]
    t = Table(filas, colWidths=col_w, repeatRows=1)
    estilo = _estilo_tabla_base(7)
    estilo.add('ALIGN', (6, 1), (6, -1), 'RIGHT')
    t.setStyle(estilo)
    elementos.append(t)

    subtitulo = f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}  |  Generado: {timezone.now().astimezone(ZONA_COLOMBIA).strftime('%d/%m/%Y %H:%M')}"
    return _build_pdf("Reporte de Devoluciones", subtitulo, elementos, orientacion='landscape')


def _pdf_compras(compras_qs, fecha_inicio, fecha_fin):
    elementos = []
    total_valor       = sum(c.valor for c in compras_qs)
    total_ordenes     = compras_qs.count()
    total_proveedores = compras_qs.values('proveedor').distinct().count()

    elementos.append(_kpi_table([
        ('Valor Total',  f'${total_valor:,.0f}', '#2ecc71'),
        ('Órdenes',      str(total_ordenes),     '#4DA8DA'),
        ('Proveedores',  str(total_proveedores), '#f1c40f'),
    ]))
    elementos.append(Spacer(1, 12))

    cabecera = [['#', 'Proveedor', 'Cant.', 'Precio Unit.', 'Total Compra', 'Estado', 'Fecha']]
    filas = cabecera
    contador = 1
    for c in compras_qs:
        det = c.detalles.first()
        filas.append([
            str(contador),
            c.proveedor.nombre_empresa,
            str(det.cantidad if det else 1),
            f'${float(det.precio_unitario if det else c.valor):,.0f}',
            f'${float(c.valor):,.0f}',
            c.get_estado_display(),
            c.fecha.astimezone(ZONA_COLOMBIA).strftime("%d/%m/%Y"),
        ])
        contador += 1

    if len(filas) == 1:
        filas.append(['', 'Sin compras en el período', '', '', '', '', ''])

    col_w = [1*cm, 5*cm, 2*cm, 3*cm, 3*cm, 2.5*cm, 2.5*cm]
    t = Table(filas, colWidths=col_w, repeatRows=1)
    estilo = _estilo_tabla_base(7)
    estilo.add('ALIGN', (2, 1), (4, -1), 'RIGHT')
    t.setStyle(estilo)
    elementos.append(t)

    subtitulo = f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}  |  Generado: {timezone.now().astimezone(ZONA_COLOMBIA).strftime('%d/%m/%Y %H:%M')}"
    return _build_pdf("Reporte de Compras", subtitulo, elementos, orientacion='landscape')


# ─────────────────────────────────────────────
#  HELPERS EXCEL
# ─────────────────────────────────────────────
def _xls_fills():
    return {
        'marca':  PatternFill('solid', fgColor=XLS_ROJO_MARCA),
        'header': PatternFill('solid', fgColor=XLS_ROJO_HEADER),
        'suave':  PatternFill('solid', fgColor=XLS_ROJO_SUAVE),
        'blanco': PatternFill('solid', fgColor=XLS_BLANCO),
        'kpi_bg': PatternFill('solid', fgColor=XLS_ROJO_KPI_BG),
    }


def _xls_borde_suave():
    lado = Side(style='thin', color=XLS_ROJO_BORDE)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _xls_encabezado(ws, titulo, subtitulo, meta_lineas, num_cols):
    ultima_col = get_column_letter(max(num_cols, 4))
    ws.merge_cells(f'A1:{ultima_col}1')
    c = ws['A1']
    c.value = 'CYS Ltda.'
    c.font = Font(name=XLS_FONT, size=18, bold=True, color=XLS_ROJO_MARCA)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f'A2:{ultima_col}2')
    c = ws['A2']
    c.value = titulo
    c.font = Font(name=XLS_FONT, size=13, bold=True, color=XLS_ROJO_HEADER)
    ws.row_dimensions[2].height = 20

    ws.merge_cells(f'A3:{ultima_col}3')
    c = ws['A3']
    c.value = subtitulo
    c.font = Font(name=XLS_FONT, size=9, italic=True, color=XLS_GRIS_TEXTO)
    ws.row_dimensions[3].height = 16

    fila = 4
    if meta_lineas:
        for etiqueta, valor in meta_lineas:
            ws.cell(row=fila, column=1, value=etiqueta).font = Font(name=XLS_FONT, size=9, bold=True, color=XLS_GRIS_TEXTO)
            ws.cell(row=fila, column=2, value=valor).font = Font(name=XLS_FONT, size=9, color=XLS_GRIS_TEXTO)
            fila += 1
    return fila + 1


def _xls_kpis(ws, fila, kpis, num_cols):
    fills = _xls_fills()
    borde = _xls_borde_suave()
    col = 1
    for label, valor, color_hex in kpis:
        cv = ws.cell(row=fila, column=col, value=valor)
        cv.font = Font(name=XLS_FONT, size=14, bold=True, color=color_hex)
        cv.alignment = Alignment(horizontal='center', vertical='center')
        cv.fill = fills['kpi_bg']
        cv.border = borde

        cl = ws.cell(row=fila + 1, column=col, value=label.upper())
        cl.font = Font(name=XLS_FONT, size=8, bold=True, color=XLS_GRIS_TEXTO)
        cl.alignment = Alignment(horizontal='center', vertical='center')
        cl.fill = fills['kpi_bg']
        cl.border = borde
        col += 1

    ws.row_dimensions[fila].height = 24
    ws.row_dimensions[fila + 1].height = 16
    return fila + 3


def _xls_tabla(ws, fila, headers, filas, col_widths, currency_cols=None, right_cols=None):
    fills = _xls_fills()
    borde = _xls_borde_suave()
    currency_cols = currency_cols or set()
    right_cols = right_cols or set()

    for j, texto in enumerate(headers, start=1):
        c = ws.cell(row=fila, column=j, value=texto)
        c.font = Font(name=XLS_FONT, size=10, bold=True, color=XLS_BLANCO)
        c.fill = fills['header']
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = borde
    ws.row_dimensions[fila].height = 20

    for i, datos_fila in enumerate(filas):
        r = fila + 1 + i
        fondo = fills['suave'] if i % 2 == 1 else fills['blanco']
        for j, valor in enumerate(datos_fila, start=1):
            c = ws.cell(row=r, column=j, value=valor)
            c.fill = fondo
            c.border = borde
            c.font = Font(name=XLS_FONT, size=9, color=XLS_GRIS_TEXTO)
            if (j - 1) in currency_cols and isinstance(valor, (int, float)):
                c.number_format = '"$"#,##0'
                c.alignment = Alignment(horizontal='right')
            elif (j - 1) in right_cols:
                c.alignment = Alignment(horizontal='right')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[r].height = 16

    for j, ancho in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho

    return fila + len(filas) + 2


# ── GENERADORES EXCEL ────────────────────────────────────────────────────────
def _xlsx_ventas(wb, ventas_qs, fecha_inicio, fecha_fin):
    ws = wb.active
    ws.title = 'Historial de Ventas'
    ws.sheet_view.showGridLines = False

    total_v = sum(v.total_venta for v in ventas_qs)
    total_u = sum(det.cantidad for v in ventas_qs for det in v.detalleventa_set.all())

    fila = _xls_encabezado(ws, "Historial de Ventas", f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}", [], 6)
    fila = _xls_kpis(ws, fila, [
        ('Total Ventas', float(total_v), XLS_VERDE_OK),
        ('Unidades', total_u, XLS_ROJO_HEADER),
        ('Órdenes', ventas_qs.count(), XLS_ROJO_MARCA),
    ], 6)

    headers = ['#', 'Fecha', 'Venta ID', 'Producto', 'Cantidad', 'Precio Unit.', 'Total']
    filas = []
    cont = 1
    for v in ventas_qs:
        for det in v.detalleventa_set.all():
            filas.append([
                cont,
                v.fecha.astimezone(ZONA_COLOMBIA).strftime("%Y-%m-%d %H:%M"),
                f"#{v.codigo_venta}",
                det.producto.nombre if det.producto else '—',
                det.cantidad,
                float(det.precio_unitario),
                float(det.subtotal),
            ])
            cont += 1

    _xls_tabla(ws, fila, headers, filas, [6, 18, 12, 32, 10, 14, 14], currency_cols={5, 6}, right_cols={4})


def _xlsx_compras(wb, compras_qs, fecha_inicio, fecha_fin):
    ws = wb.active
    ws.title = 'Compras'
    ws.sheet_view.showGridLines = False

    total_valor = sum(c.valor for c in compras_qs)
    fila = _xls_encabezado(ws, "Reporte de Compras", f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}", [], 6)
    fila = _xls_kpis(ws, fila, [
        ('Total Compras', float(total_valor), XLS_VERDE_OK),
        ('Órdenes', compras_qs.count(), XLS_ROJO_MARCA),
    ], 6)

    headers = ['#', 'Proveedor', 'Cant.', 'Precio Unit.', 'Valor Total', 'Estado', 'Fecha']
    filas = []
    for i, c in enumerate(compras_qs, 1):
        det = c.detalles.first()
        filas.append([
            i,
            c.proveedor.nombre_empresa,
            det.cantidad if det else 1,
            float(det.precio_unitario if det else c.valor),
            float(c.valor),
            c.get_estado_display(),
            c.fecha.astimezone(ZONA_COLOMBIA).strftime("%d/%m/%Y"),
        ])
    _xls_tabla(ws, fila, headers, filas, [6, 28, 10, 14, 15, 14, 14], currency_cols={3, 4}, right_cols={2})


def _xlsx_proveedores(wb, proveedores_list):
    ws = wb.active
    ws.title = 'Proveedores'
    ws.sheet_view.showGridLines = False

    fila = _xls_encabezado(ws, "Reporte de Proveedores", "Proveedores registrados en CYS", [], 5)
    fila = _xls_kpis(ws, fila, [('Total Proveedores', len(proveedores_list), XLS_ROJO_MARCA)], 5)

    headers = ['#', 'NIT', 'Empresa', 'Teléfono', 'Correo', 'Estado']
    filas = [[i, p.nit_proveedores, p.nombre_empresa, p.telefono or '—', p.correo, p.get_estado_display()] for i, p in enumerate(proveedores_list, 1)]
    _xls_tabla(ws, fila, headers, filas, [6, 16, 30, 16, 30, 14])


def _xlsx_devoluciones(wb, devoluciones_qs, fecha_inicio, fecha_fin):
    ws = wb.active
    ws.title = 'Devoluciones'
    ws.sheet_view.showGridLines = False

    total_valor = sum(d.total_devuelto for d in devoluciones_qs)
    fila = _xls_encabezado(ws, "Reporte de Devoluciones", f"Período: {fecha_inicio or 'Todo'} → {fecha_fin or 'Todo'}", [], 6)
    fila = _xls_kpis(ws, fila, [
        ('Devoluciones', devoluciones_qs.count(), XLS_ROJO_MARCA),
        ('Total Devuelto', float(total_valor), XLS_ROJO_OUT),
    ], 6)

    headers = ['#', 'Fecha', 'Venta Asoc.', 'Motivo', 'Tipo', 'Estado', 'Total']
    filas = [[
        i,
        d.fecha.astimezone(ZONA_COLOMBIA).strftime("%d/%m/%Y %H:%M"),
        f"Venta #{d.venta.codigo_venta}" if d.venta else '—',
        d.motivo,
        d.tipo_devolucion,
        d.estado.capitalize(),
        float(d.total_devuelto),
    ] for i, d in enumerate(devoluciones_qs, 1)]
    _xls_tabla(ws, fila, headers, filas, [6, 18, 14, 25, 18, 14, 15], currency_cols={6})


# ─────────────────────────────────────────────
#  VISTA PRINCIPAL: INDEX REPORTES
# ─────────────────────────────────────────────
def index_reportes(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin    = request.GET.get('fecha_fin')
    per_page     = int(request.GET.get('per_page', 10) or 10)

    # Ventas
    ventas_qs = Venta.objects.prefetch_related('detalleventa_set__producto').all().order_by('-fecha')
    if fecha_inicio:
        ventas_qs = ventas_qs.filter(fecha__date__gte=fecha_inicio)
    if fecha_fin:
        ventas_qs = ventas_qs.filter(fecha__date__lte=fecha_fin)

    # Devoluciones
    devoluciones_qs = Devolucion.objects.select_related('venta').all().order_by('-fecha')
    if fecha_inicio:
        devoluciones_qs = devoluciones_qs.filter(fecha__date__gte=fecha_inicio)
    if fecha_fin:
        devoluciones_qs = devoluciones_qs.filter(fecha__date__lte=fecha_fin)

    # Compras
    compras_qs = Compra.objects.select_related('proveedor').prefetch_related('detalles').all().order_by('-fecha')
    if fecha_inicio:
        compras_qs = compras_qs.filter(fecha__date__gte=fecha_inicio)
    if fecha_fin:
        compras_qs = compras_qs.filter(fecha__date__lte=fecha_fin)

    # Paginación Ventas
    paginator   = Paginator(ventas_qs, per_page)
    page_number = request.GET.get('page', 1)
    page_obj    = paginator.get_page(page_number)

    # Proveedores reales de la BD
    proveedores = Proveedor.objects.all().order_by('nombre_empresa')

    # ─────────────────────────────────────────────
    #  EXPORTACIONES (PDF / EXCEL)
    # ─────────────────────────────────────────────
    export_format = request.GET.get('export')
    if export_format in ('excel', 'pdf'):
        tipo = request.GET.get('tipo', 'ventas')
        if tipo not in TIPOS_VALIDOS:
            return HttpResponse(f"Tipo de reporte '{tipo}' no válido.", status=400)

        # Exportar Excel
        if export_format == 'excel':
            wb = Workbook()
            if tipo == 'ventas':
                _xlsx_ventas(wb, ventas_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_ventas_{timezone.now().strftime('%Y%m%d')}.xlsx"
            elif tipo == 'proveedores':
                _xlsx_proveedores(wb, list(proveedores))
                nombre = f"reporte_proveedores_{timezone.now().strftime('%Y%m%d')}.xlsx"
            elif tipo == 'devoluciones':
                _xlsx_devoluciones(wb, devoluciones_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_devoluciones_{timezone.now().strftime('%Y%m%d')}.xlsx"
            elif tipo == 'compras':
                _xlsx_compras(wb, compras_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_compras_{timezone.now().strftime('%Y%m%d')}.xlsx"
            else:
                _xlsx_ventas(wb, ventas_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_{tipo}_{timezone.now().strftime('%Y%m%d')}.xlsx"

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{nombre}"'
            return response

        # Exportar PDF
        elif export_format == 'pdf':
            hoy_pdf = timezone.now().astimezone(ZONA_COLOMBIA).date()
            if tipo == 'ventas':
                buffer = _pdf_ventas(ventas_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_ventas_{timezone.now().strftime('%Y%m%d')}.pdf"
            elif tipo == 'proveedores':
                buffer = _pdf_proveedores(list(proveedores))
                nombre = f"reporte_proveedores_{timezone.now().strftime('%Y%m%d')}.pdf"
            elif tipo == 'devoluciones':
                buffer = _pdf_devoluciones(devoluciones_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_devoluciones_{timezone.now().strftime('%Y%m%d')}.pdf"
            elif tipo == 'compras':
                buffer = _pdf_compras(compras_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_compras_{timezone.now().strftime('%Y%m%d')}.pdf"
            else:
                buffer = _pdf_ventas(ventas_qs, fecha_inicio, fecha_fin)
                nombre = f"reporte_{tipo}_{timezone.now().strftime('%Y%m%d')}.pdf"

            response = HttpResponse(buffer, content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="{nombre}"'
            return response

    # ─────────────────────────────────────────────
    #  CÁLCULO DE KPIS Y DATOS PARA EL TEMPLATE
    # ─────────────────────────────────────────────
    hoy = timezone.now().astimezone(ZONA_COLOMBIA).date()

    total_ventas = sum(v.total_venta for v in ventas_qs)
    total_productos = sum(det.cantidad for v in ventas_qs for det in v.detalleventa_set.all())
    total_ordenes_ventas = ventas_qs.count()

    ventas_hoy = Venta.objects.filter(fecha__date=hoy).prefetch_related('detalleventa_set__producto')
    ingresos_hoy = sum(v.total_venta for v in ventas_hoy)

    # Top productos vendidos hoy
    detalles_hoy = DetalleVenta.objects.filter(venta__fecha__date=hoy).select_related('producto')
    top_h = {}
    for det in detalles_hoy:
        nombre = det.producto.nombre if det.producto else 'Sin producto'
        if nombre not in top_h:
            top_h[nombre] = {'cantidad': 0, 'subtotal': 0.0}
        top_h[nombre]['cantidad'] += det.cantidad
        top_h[nombre]['subtotal'] += float(det.subtotal)
    top_productos_hoy = sorted(top_h.items(), key=lambda x: x[1]['subtotal'], reverse=True)[:5]

    # Devoluciones
    total_devoluciones = sum(d.total_devuelto for d in devoluciones_qs)
    cantidad_devoluciones = devoluciones_qs.count()
    devoluciones_aprobadas = sum(1 for d in devoluciones_qs if d.estado == 'aprobada')

    # Compras
    total_compras = sum(c.valor for c in compras_qs)
    total_ordenes_compras = compras_qs.count()

    context = {
        'ventas':                ventas_qs,
        'page_obj':              page_obj,
        'paginator':             paginator,
        'total_ventas':          total_ventas,
        'total_productos':       total_productos,
        'total_clientes':        total_ordenes_ventas,
        'proveedores':           proveedores,
        'fecha_inicio':          fecha_inicio or '',
        'fecha_fin':             fecha_fin or '',
        'per_page':              per_page,
        'hoy':                   hoy,
        'ventas_hoy':            ventas_hoy,
        'ingresos_hoy':          ingresos_hoy,
        'top_productos_hoy':     top_productos_hoy,
        'total_devoluciones':    total_devoluciones,
        'cantidad_devoluciones': cantidad_devoluciones,
        'devoluciones_aprobadas': devoluciones_aprobadas,
        'total_compras':         total_compras,
        'total_ordenes_compras': total_ordenes_compras,
        'breadcrumb_items': [
            {'nombre': 'Inicio', 'url': reverse('principal')},
            {'nombre': 'Reportes', 'url': None},
        ],
    }

    return render(request, 'reportes/reportes.html', context)


# ─────────────────────────────────────────────
#  VISTA AJAX — RESUMEN DIARIO
# ─────────────────────────────────────────────
def resumen_diario_ajax(request):
    """
    Vista AJAX usada por el modal 'Resumen Diario' para consultar
    los datos de ventas de un día específico.
    """
    fecha_param = request.GET.get('fecha')
    if fecha_param:
        try:
            fecha_consulta = timezone.datetime.strptime(fecha_param, '%Y-%m-%d').date()
        except ValueError:
            fecha_consulta = timezone.now().astimezone(ZONA_COLOMBIA).date()
    else:
        fecha_consulta = timezone.now().astimezone(ZONA_COLOMBIA).date()

    ventas_dia = Venta.objects.filter(fecha__date=fecha_consulta).prefetch_related('detalleventa_set__producto')
    ingresos = sum(v.total_venta for v in ventas_dia)
    unidades = sum(det.cantidad for v in ventas_dia for det in v.detalleventa_set.all())

    data = {
        'success': True,
        'fecha': fecha_consulta.strftime('%d/%m/%Y'),
        'total_ventas': len(ventas_dia),
        'ingresos': float(ingresos),
        'unidades_vendidas': unidades,
    }
    return JsonResponse(data)