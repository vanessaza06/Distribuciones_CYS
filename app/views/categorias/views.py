import io
from datetime import datetime

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import Categoria


@login_required
def categorias_lista(request):
    # Lista plana: ya no hay jerarquía padre/hijo
    categorias = Categoria.objects.prefetch_related('productos').all()

    context = {
        'categorias': categorias,
        'total_activas': categorias.filter(activo=True).count(),
        'total_inactivas': categorias.filter(activo=False).count(),
        'breadcrumb_items': [
            {'nombre': 'Categorías', 'url': None},
        ],
    }
    return render(request, 'categorias/categorias.html', context)


@login_required
def categoria_crear(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        codigo = request.POST.get('codigo', '').strip().upper()
        descripcion = request.POST.get('descripcion', '').strip()

        if not nombre or not codigo:
            messages.error(request, '⚠️ Nombre y código son obligatorios.')
            return redirect('categorias_lista')

        if Categoria.objects.filter(codigo__iexact=codigo).exists():
            messages.error(request, f'⚠️ Ya existe una categoría con el código "{codigo}".')
            return redirect('categorias_lista')

        Categoria.objects.create(
            nombre=nombre,
            codigo=codigo,
            descripcion=descripcion,
        )
        messages.success(request, f'✅ Categoría "{nombre}" creada con éxito.')

    return redirect('categorias_lista')


@login_required
def categoria_editar(request, pk):
    categoria = get_object_or_404(Categoria, pk=pk)

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        codigo = request.POST.get('codigo', '').strip().upper()
        descripcion = request.POST.get('descripcion', '').strip()

        if not nombre or not codigo:
            messages.error(request, '⚠️ Nombre y código son obligatorios.')
            return redirect('categorias_lista')

        if Categoria.objects.filter(codigo__iexact=codigo).exclude(pk=pk).exists():
            messages.error(request, f'⚠️ Ya existe otra categoría con el código "{codigo}".')
            return redirect('categorias_lista')

        categoria.nombre = nombre
        categoria.codigo = codigo
        categoria.descripcion = descripcion
        categoria.save()

        messages.success(request, f'✅ Categoría "{nombre}" actualizada correctamente.')

    return redirect('categorias_lista')


@login_required
def categoria_toggle_activo(request, pk):
    """Las categorías no se eliminan: solo se activan o desactivan."""
    categoria = get_object_or_404(Categoria, pk=pk)

    if request.method == 'POST':
        categoria.activo = not categoria.activo
        categoria.save(update_fields=['activo'])

        if categoria.activo:
            messages.success(request, f'✅ Categoría "{categoria.nombre}" activada.')
        else:
            messages.warning(
                request,
                f'⚠️ Categoría "{categoria.nombre}" desactivada. No se elimina para conservar el histórico de productos.'
            )

    return redirect('categorias_lista')


# ============================================================
# EXPORTACIONES
# ============================================================

def _categorias_para_exportar():
    return Categoria.objects.prefetch_related('productos').order_by('nombre')


@login_required
def categoria_exportar_excel(request):
    categorias = _categorias_para_exportar()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Categorías'

    # Título
    ws.merge_cells('A1:E1')
    ws['A1'] = 'Listado de Categorías - Licorera'
    ws['A1'].font = Font(size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='021A3A')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:E2')
    ws['A2'] = f'Generado: {datetime.now():%d/%m/%Y %H:%M}'
    ws['A2'].alignment = Alignment(horizontal='center')
    ws['A2'].font = Font(italic=True, color='666666')

    # Encabezados
    headers = ['Nombre', 'Código', 'Descripción', 'N° Productos', 'Estado']
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F6FD6')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border

    # Datos
    row = 5
    for cat in categorias:
        valores = [
            cat.nombre,
            cat.codigo,
            cat.descripcion or '',
            cat.productos.count(),
            'Activa' if cat.activo else 'Inactiva',
        ]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = border
            if col in (2, 4, 5):
                c.alignment = Alignment(horizontal='center')
        row += 1

    # Anchos de columna
    for i, w in enumerate([28, 14, 50, 16, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A5'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename="categorias_{datetime.now():%Y%m%d_%H%M}.xlsx"'
    )
    wb.save(response)
    return response


@login_required
def categoria_exportar_pdf(request):
    categorias = _categorias_para_exportar()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title='Listado de Categorías',
    )

    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('titulo', parent=styles['Title'], textColor=colors.HexColor('#021A3A'))
    sub = ParagraphStyle('sub', parent=styles['Normal'], textColor=colors.grey, alignment=1)
    celda = ParagraphStyle('celda', parent=styles['Normal'], fontSize=9, leading=11)

    elementos = [
        Paragraph('Listado de Categorías - Licorera', titulo),
        Paragraph(f'Generado: {datetime.now():%d/%m/%Y %H:%M}', sub),
        Spacer(1, 0.6 * cm),
    ]

    data = [['Nombre', 'Código', 'Descripción', 'Productos', 'Estado']]
    for cat in categorias:
        data.append([
            Paragraph(cat.nombre, celda),
            cat.codigo,
            Paragraph(cat.descripcion or '-', celda),
            str(cat.productos.count()),
            'Activa' if cat.activo else 'Inactiva',
        ])

    tabla = Table(
        data, repeatRows=1,
        colWidths=[4 * cm, 2.2 * cm, 6.3 * cm, 2.2 * cm, 2.3 * cm],
    )
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#021A3A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('ALIGN', (3, 0), (4, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6FC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elementos.append(tabla)

    doc.build(elementos)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="categorias_{datetime.now():%Y%m%d_%H%M}.pdf"'
    )
    return response