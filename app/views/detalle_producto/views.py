from datetime import datetime
from io import BytesIO

from django.db.models import Sum, Q
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse, Http404
from django.utils.html import escape
from app.models import Categoria, DetalleProducto, Producto, Marca
from app.forms import DetalleProductoForm
from django.urls import reverse


@login_required
def detalle_producto_lista(request):
    productos_qs = Producto.objects.select_related('categoria__subcategoria')

    con_codigo = productos_qs.filter(detalles__codigo_barras__isnull=False).exclude(detalles__codigo_barras='').distinct().count()

    # Una sola pasada: stock anotado + relaciones precargadas (antes ~6800 consultas)
    productos = list(
        productos_qs
        .annotate(stock_total_anotado=Sum('lotes__stock_actual'))
        .prefetch_related('detalles', 'presentaciones__lotes')
    )
    sin_codigo = len(productos) - con_codigo

    # Se evalúa una vez: la plantilla recorre las categorías dentro de cada modal
    categorias = list(
        Categoria.objects.filter(subcategoria__isnull=True).prefetch_related('subcategorias')
    )

    context = {
        'productos': productos,
        'categorias': categorias,  # solo principales
        'con_codigo': con_codigo,
        'sin_codigo': sin_codigo,
        'marcas': Marca.objects.all(),
        'breadcrumb_items': [
            {'nombre': 'Inventario', 'url': None},
            {'nombre': 'Detalle de Producto', 'url': None},
        ],
    }
    return render(request, 'detalle_producto/detalle_producto.html', context)


def _filas_exportacion(request):
    """Devuelve (encabezados, filas) respetando el texto del buscador (?q=)."""
    qs = (
        Producto.objects
        .select_related('categoria')
        .annotate(stock_total_anotado=Sum('lotes__stock_actual'))
        .prefetch_related('detalles', 'presentaciones__lotes')
        .order_by('nombre')
    )

    q = request.GET.get('q', '').strip()
    if q:
        # Subconsulta para no duplicar filas ni alterar la suma del stock
        qs = qs.filter(
            Q(nombre__icontains=q) |
            Q(pk__in=DetalleProducto.objects.filter(codigo_barras__icontains=q).values('producto_id'))
        )

    encabezados = [
        'Producto', 'Categoría', 'Código de barras', 'Stock total',
        'Costo unit. (último lote)', 'Presentaciones', 'Estado',
    ]

    filas = []
    for p in qs:
        detalle = p.detalle_principal
        lote = p.lote_referencia
        costo = lote.costo_unitario if lote and lote.costo_unitario else ''
        presentaciones = ' | '.join(
            f'{pr.nombre} (x{pr.cantidad}) ${pr.precio_venta:,.0f}'
            for pr in p.presentaciones.all()
        )
        filas.append([
            p.nombre,
            p.categoria.nombre if p.categoria else '',
            detalle.codigo_barras if detalle and detalle.codigo_barras else '',
            p.stock_total_anotado or 0,
            costo,
            presentaciones,
            'Activo' if p.activo else 'Inactivo',
        ])
    return encabezados, filas


def _costo_texto(valor):
    return f'${valor:,.0f}' if valor != '' else '—'


@login_required
def detalle_producto_exportar(request, formato):
    """Exporta el listado de productos con su detalle: excel, pdf o imprimir."""
    if formato not in ('excel', 'pdf', 'imprimir'):
        raise Http404('Formato no soportado.')

    encabezados, filas = _filas_exportacion(request)
    ahora = datetime.now()
    nombre_archivo = f'detalle_productos_{ahora.strftime("%Y%m%d_%H%M")}'
    fecha_texto = ahora.strftime('%d/%m/%Y %H:%M')

    # ───────────── EXCEL ─────────────
    if formato == 'excel':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = 'Detalle de Productos'
        ws.append(encabezados)

        relleno = PatternFill('solid', fgColor='0B2A4A')
        for celda in ws[1]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = relleno
            celda.alignment = Alignment(horizontal='center', vertical='center')

        for fila in filas:
            ws.append(fila)

        for fila in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            for celda in fila:
                celda.number_format = '"$"#,##0'

        for i, col in enumerate(ws.columns, start=1):
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 3, 60)

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.xlsx"'
        wb.save(response)
        return response

    # ───────────── PDF ─────────────
    if formato == 'pdf':
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=landscape(A4),
            leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm,
            title='Detalle de Productos',
        )
        estilos = getSampleStyleSheet()
        celda = ParagraphStyle('celda', parent=estilos['Normal'], fontSize=7, leading=9)
        celda_head = ParagraphStyle('celda_head', parent=celda, textColor=colors.white, fontName='Helvetica-Bold')

        datos = [[Paragraph(h, celda_head) for h in encabezados]]
        for f in filas:
            datos.append([
                Paragraph(escape(str(f[0])), celda),
                Paragraph(escape(str(f[1])), celda),
                Paragraph(escape(str(f[2])), celda),
                Paragraph(str(f[3]), celda),
                Paragraph(_costo_texto(f[4]), celda),
                Paragraph(escape(str(f[5])), celda),
                Paragraph(f[6], celda),
            ])

        tabla = Table(
            datos, repeatRows=1,
            colWidths=[5 * cm, 3.2 * cm, 3.2 * cm, 2 * cm, 2.8 * cm, 7.5 * cm, 2 * cm],
        )
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0B2A4A')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6FB')]),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#C5D2E3')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        doc.build([
            Paragraph('Detalle de Productos — CYS Ltda', estilos['Title']),
            Paragraph(f'Generado: {fecha_texto} · Total: {len(filas)} productos', estilos['Normal']),
            Spacer(1, 0.4 * cm),
            tabla,
        ])

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.pdf"'
        return response

    # ───────────── IMPRIMIR ─────────────
    thead = ''.join(f'<th>{escape(h)}</th>' for h in encabezados)
    tbody = ''.join(
        '<tr>'
        f'<td>{escape(str(f[0]))}</td>'
        f'<td>{escape(str(f[1]))}</td>'
        f'<td>{escape(str(f[2]))}</td>'
        f'<td>{f[3]}</td>'
        f'<td>{_costo_texto(f[4])}</td>'
        f'<td>{escape(str(f[5]))}</td>'
        f'<td>{f[6]}</td>'
        '</tr>'
        for f in filas
    )
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>Detalle de Productos</title>
<style>
  body {{ font-family: Arial, sans-serif; margin: 20px; color: #111; }}
  h1 {{ font-size: 18px; margin: 0 0 4px; }}
  p {{ font-size: 11px; margin: 0 0 12px; color: #555; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 10px; }}
  th {{ background: #0B2A4A; color: #fff; padding: 5px; text-align: left; }}
  td {{ border-bottom: 1px solid #ccd6e3; padding: 4px 5px; vertical-align: top; }}
  tr:nth-child(even) td {{ background: #f2f6fb; }}
  thead {{ display: table-header-group; }}
  tr {{ page-break-inside: avoid; }}
  @page {{ size: A4 landscape; margin: 10mm; }}
</style>
</head>
<body>
  <h1>Detalle de Productos — CYS Ltda</h1>
  <p>Generado: {fecha_texto} · Total: {len(filas)} productos</p>
  <table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table>
  <script>window.onload = function () {{ window.print(); }};</script>
</body>
</html>"""
    return HttpResponse(html)


@login_required
def detalle_producto_crear_rapido(request, producto_pk):
    producto = get_object_or_404(Producto, pk=producto_pk)
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    if DetalleProducto.objects.filter(producto=producto).exists():
        return JsonResponse({'ok': False, 'error': 'Este producto ya tiene un detalle.'}, status=400)

    form = DetalleProductoForm(request.POST)
    if not form.is_valid():
        primer_error = next(iter(form.errors.values()))[0]
        return JsonResponse({'ok': False, 'error': primer_error}, status=400)

    detalle = form.save(commit=False)
    detalle.producto = producto
    detalle.save()

    return JsonResponse({
        'ok': True,
        'detalle_pk': detalle.pk,
        'codigo_barras': detalle.codigo_barras,
        'guardar_url': reverse('guardar_codigo', args=[detalle.pk]),
    })


@login_required
def detalle_producto_crear(request, producto_pk):
    producto = get_object_or_404(Producto, pk=producto_pk)

    if request.method == 'POST':
        if DetalleProducto.objects.filter(producto=producto).exists():
            messages.error(request, 'Este producto ya tiene un detalle registrado.')
            return redirect('producto_detalle', pk=producto_pk)

        form = DetalleProductoForm(request.POST)
        if form.is_valid():
            detalle = form.save(commit=False)
            detalle.producto = producto
            detalle.save()
            messages.success(request, 'Detalle de producto registrado.')
        else:
            for campo, errores in form.errors.items():
                for error in errores:
                    messages.error(request, f'{campo}: {error}')

    return redirect('producto_detalle', pk=producto_pk)


@login_required
def detalle_producto_editar(request, pk):
    detalle = get_object_or_404(DetalleProducto, pk=pk)

    if request.method == 'POST':
        form = DetalleProductoForm(request.POST, instance=detalle)
        if form.is_valid():
            form.save()
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'ok': True})
            messages.success(request, 'Detalle de producto actualizado.')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'ok': False, 'errores': form.errors}, status=400)
            for campo, errores in form.errors.items():
                for error in errores:
                    messages.error(request, f'{campo}: {error}')

    return redirect('detalle_producto_lista')


@login_required
def guardar_codigo(request, pk):
    """Endpoint rápido solo para el código de barras (ej: desde escaneo con cámara)."""
    detalle = get_object_or_404(DetalleProducto, pk=pk)
    if request.method == 'POST':
        codigo = request.POST.get('codigo_barras', '').strip()
        if not codigo:
            return JsonResponse({'ok': False, 'error': 'Código vacío.'}, status=400)
        if DetalleProducto.objects.filter(codigo_barras=codigo).exclude(pk=detalle.pk).exists():
            return JsonResponse({'ok': False, 'error': 'Ese código ya existe.'}, status=400)

        detalle.codigo_barras = codigo
        detalle.save(update_fields=['codigo_barras'])
        return JsonResponse({'ok': True, 'codigo_barras': detalle.codigo_barras})

    return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)