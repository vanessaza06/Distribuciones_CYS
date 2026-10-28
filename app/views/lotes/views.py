from decimal import Decimal, InvalidOperation
from io import BytesIO
import json

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import FieldDoesNotExist
from django.db import models
from django.db.models import Sum
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape

from app.models import Producto, Lote, Bodega, PresentacionProducto
from app.forms import LoteForm

# Nombre del campo de fecha de ingreso en el modelo Lote.
# Si tu modelo lo llama distinto, cámbialo aquí.
CAMPO_FECHA_LOTE = 'fecha_ingreso'


def _stats_lotes(hoy):
    """Devuelve (unidades ingresadas hoy, lotes registrados en el mes)."""
    try:
        campo = Lote._meta.get_field(CAMPO_FECHA_LOTE)
    except FieldDoesNotExist:
        return 0, 0

    if isinstance(campo, models.DateTimeField):
        filtro_hoy = {f'{CAMPO_FECHA_LOTE}__date': hoy}
    else:
        filtro_hoy = {CAMPO_FECHA_LOTE: hoy}

    ingresos_hoy = (
        Lote.objects.filter(**filtro_hoy)
        .aggregate(total=Sum('cantidad_inicial'))['total'] or 0
    )
    lotes_mes = Lote.objects.filter(
        **{
            f'{CAMPO_FECHA_LOTE}__year': hoy.year,
            f'{CAMPO_FECHA_LOTE}__month': hoy.month,
        }
    ).count()
    return ingresos_hoy, lotes_mes


@login_required
def lote_list(request):
    """Página 'Lotes': formulario de registro + resumen.

    Si llega ?presentacion=<pk> (por ejemplo desde el botón "Crear Lote"
    de Bodega), esa presentación queda preseleccionada en el formulario.
    """
    hoy = timezone.localdate()

    # Si no hay ninguna bodega, se crea una por defecto para no bloquear el lote.
    if not Bodega.objects.exists():
        Bodega.objects.get_or_create(nombre='Bodega Principal')

    productos = Producto.objects.prefetch_related('presentaciones')
    bodegas = Bodega.objects.all()
    hay_presentaciones = PresentacionProducto.objects.exists()
    hay_bodegas = bodegas.exists()

    presentacion_preseleccionada = request.GET.get('presentacion', '')

    ingresos_hoy, ordenes_mes = _stats_lotes(hoy)

    top = (
        Lote.objects
        .values('presentacion__producto__nombre')
        .annotate(total=Sum('stock_actual'))
        .order_by('-total')[:5]
    )
    labels = [t['presentacion__producto__nombre'] for t in top]
    data = [int(t['total'] or 0) for t in top]

    context = {
        'productos': productos,
        'bodegas': bodegas,
        'hay_presentaciones': hay_presentaciones,
        'hay_bodegas': hay_bodegas,
        'presentacion_preseleccionada': presentacion_preseleccionada,
        'ingresos_hoy': ingresos_hoy,
        'ordenes_mes': ordenes_mes,
        'proveedores_labels': json.dumps(labels),
        'proveedores_data': json.dumps(data),
        'breadcrumb_items': [
            {'nombre': 'Inventario', 'url': None},
            {'nombre': 'Lotes', 'url': reverse('lote_list')},
        ],
    }
    return render(request, 'lotes/lotes.html', context)


@login_required
def gestion_stock(request):
    productos = Producto.objects.select_related('categoria').prefetch_related('presentaciones__lotes')
    lotes_activos = list(Lote.objects.select_related('presentacion__producto', 'bodega'))
    lotes_por_vencer = [l for l in lotes_activos if l.stock_actual > 0 and l.proximo_a_vencer]
    lotes_vencidos = [l for l in lotes_activos if l.stock_actual > 0 and l.esta_vencido]

    context = {
        'productos': productos,
        'lotes_activos': lotes_activos,
        'lotes_por_vencer': lotes_por_vencer,
        'lotes_vencidos': lotes_vencidos,
        'breadcrumb_items': [
            {'nombre': 'Inventario', 'url': None},
            {'nombre': 'Stock & Productos', 'url': reverse('gestion_stock')},
        ],
    }
    return render(request, 'lotes/gestion.html', context)


@login_required
def lote_create(request):
    if request.method == 'POST':
        datos = request.POST.copy()

        # Si no llegó bodega, se usa la primera disponible.
        if not datos.get('bodega'):
            bodega_default = Bodega.objects.first()
            if bodega_default:
                datos['bodega'] = bodega_default.pk

        form = LoteForm(datos, request.FILES)
        if form.is_valid():
            lote = form.save(commit=False)
            lote.stock_actual = lote.cantidad_inicial
            lote.costo_total = lote.costo_unitario * lote.cantidad_inicial
            lote.save()
            messages.success(request, f'Lote {lote.numero_lote} registrado.')
        else:
            for campo, errores in form.errors.items():
                for error in errores:
                    messages.error(request, f'{campo}: {error}')
    return redirect('lote_list')


@login_required
def lote_update(request, numero_lote):
    lote = get_object_or_404(Lote, numero_lote=numero_lote)
    cantidad_inicial_original = lote.cantidad_inicial

    if request.method == 'POST':
        form = LoteForm(request.POST, request.FILES, instance=lote)
        if form.is_valid():
            lote_actualizado = form.save(commit=False)
            # La cantidad inicial no se edita aquí: representa el ingreso
            # original del lote. Para cambiar el stock disponible se usa
            # "Ajustar stock".
            lote_actualizado.cantidad_inicial = cantidad_inicial_original
            lote_actualizado.costo_total = lote_actualizado.costo_unitario * cantidad_inicial_original
            lote_actualizado.save()
            messages.success(request, f'Lote {lote.numero_lote} actualizado.')
            return redirect('lote_list')
        else:
            for campo, errores in form.errors.items():
                for error in errores:
                    messages.error(request, f'{campo}: {error}')
    else:
        form = LoteForm(instance=lote)

    return render(request, 'lotes/lote_form.html', {'form': form, 'lote': lote})


@login_required
def lote_detail(request, numero_lote):
    """Reutiliza el mismo formulario de edición completa del lote."""
    lote = get_object_or_404(Lote, numero_lote=numero_lote)
    form = LoteForm(instance=lote)
    return render(request, 'lotes/lote_form.html', {'form': form, 'lote': lote})


@login_required
def lote_ajustar_stock(request, numero_lote):
    lote = get_object_or_404(Lote, numero_lote=numero_lote)

    if request.method != 'POST':
        return redirect('gestion_stock')

    nuevo_stock_raw = request.POST.get('nuevo_stock', '').strip()
    costo_unitario_raw = request.POST.get('costo_unitario', '').strip()

    try:
        nuevo_stock = int(nuevo_stock_raw)
    except (TypeError, ValueError):
        messages.error(request, 'El nuevo stock debe ser un número entero.')
        return redirect('gestion_stock')

    if nuevo_stock < 0:
        messages.error(request, 'El stock no puede ser negativo.')
        return redirect('gestion_stock')

    lote.stock_actual = nuevo_stock

    if costo_unitario_raw:
        try:
            nuevo_costo = Decimal(costo_unitario_raw.replace(',', '.'))
        except (InvalidOperation, ValueError):
            messages.error(request, 'El costo unitario no es válido.')
            return redirect('gestion_stock')

        if nuevo_costo <= 0:
            messages.error(request, 'El costo unitario debe ser mayor a 0.')
            return redirect('gestion_stock')

        lote.costo_unitario = nuevo_costo
        # Se recalcula el costo total con la cantidad inicial del lote.
        lote.costo_total = nuevo_costo * lote.cantidad_inicial
        lote.save(update_fields=['stock_actual', 'costo_unitario', 'costo_total'])
    else:
        lote.save(update_fields=['stock_actual'])

    messages.success(request, 'Stock ajustado.')
    return redirect('gestion_stock')


# ─────────────────────────── EXPORTAR LOTES ───────────────────────────

ENCABEZADOS_EXPORT = [
    'Nº Lote', 'Producto', 'Presentación', 'Bodega', 'Fecha ingreso',
    'Cant. inicial', 'Stock actual', 'Costo unitario', 'Costo total', 'Estado',
]


def _estado_lote(lote):
    if lote.stock_actual > 0 and lote.esta_vencido:
        return 'Vencido'
    if lote.stock_actual > 0 and lote.proximo_a_vencer:
        return 'Por vencer'
    if lote.stock_actual <= 0:
        return 'Agotado'
    return 'Vigente'


def _fecha_lote(lote):
    valor = getattr(lote, CAMPO_FECHA_LOTE, None)
    if not valor:
        return ''
    if hasattr(valor, 'astimezone'):          # DateTimeField
        valor = timezone.localtime(valor)
    return valor.strftime('%d/%m/%Y')


def _filas_lotes():
    """Devuelve las filas (listas) de todos los lotes para exportar."""
    lotes = (
        Lote.objects
        .select_related('presentacion__producto', 'bodega')
        .order_by('numero_lote')
    )
    filas = []
    for l in lotes:
        filas.append([
            l.numero_lote,
            l.presentacion.producto.nombre,
            l.presentacion.nombre,
            l.bodega.nombre if l.bodega else '',
            _fecha_lote(l),
            l.cantidad_inicial,
            l.stock_actual,
            l.costo_unitario,
            l.costo_total,
            _estado_lote(l),
        ])
    return filas


def _nombre_archivo(extension):
    return f'lotes_{timezone.localdate():%Y%m%d}.{extension}'


def _exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lotes'

    ws.append(ENCABEZADOS_EXPORT)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0D2A4F')
        celda.alignment = Alignment(horizontal='center', vertical='center')

    for fila in _filas_lotes():
        fila[7] = float(fila[7] or 0)
        fila[8] = float(fila[8] or 0)
        ws.append(fila)

    # Formato moneda y anchos de columna
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for celda in row:
            celda.number_format = '"$"#,##0'
    for i, titulo in enumerate(ENCABEZADOS_EXPORT, start=1):
        maximo = max(
            [len(str(titulo))] +
            [len(str(c.value or '')) for c in ws[get_column_letter(i)][1:]]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(maximo + 3, 40)
    ws.freeze_panes = 'A2'

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo("xlsx")}"'
    return response


def _exportar_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm,
        title='Reporte de Lotes',
    )
    estilos = getSampleStyleSheet()

    datos = [ENCABEZADOS_EXPORT]
    for f in _filas_lotes():
        f = list(f)
        f[7] = f'${(f[7] or 0):,.0f}'
        f[8] = f'${(f[8] or 0):,.0f}'
        datos.append(f)

    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D2A4F')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (5, 1), (8, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF3FA')]),
    ]))

    elementos = [
        Paragraph('Reporte de Lotes — CYS Ltda', estilos['Title']),
        Paragraph(f'Generado el {timezone.localdate():%d/%m/%Y}', estilos['Normal']),
        Spacer(1, 0.4 * cm),
        tabla,
    ]
    doc.build(elementos)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo("pdf")}"'
    return response


def _exportar_imprimir():
    filas = _filas_lotes()

    encabezados_html = ''.join(f'<th>{escape(e)}</th>' for e in ENCABEZADOS_EXPORT)

    filas_html = ''
    for f in filas:
        filas_html += (
            '<tr>'
            f'<td>{escape(f[0])}</td>'
            f'<td>{escape(f[1])}</td>'
            f'<td>{escape(f[2])}</td>'
            f'<td>{escape(f[3])}</td>'
            f'<td>{escape(f[4])}</td>'
            f'<td class="num">{f[5]}</td>'
            f'<td class="num">{f[6]}</td>'
            f'<td class="num">${(f[7] or 0):,.0f}</td>'
            f'<td class="num">${(f[8] or 0):,.0f}</td>'
            f'<td>{escape(f[9])}</td>'
            '</tr>'
        )
    if not filas_html:
        filas_html = '<tr><td colspan="10">No hay lotes registrados.</td></tr>'

    plural = '' if len(filas) == 1 else 's'
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <title>Reporte de Lotes — CYS Ltda</title>
  <style>
    body {{ font-family: Arial, sans-serif; font-size: 12px; color: #111; margin: 24px; }}
    h1 {{ font-size: 20px; margin: 0 0 4px; }}
    .meta {{ color: #555; margin-bottom: 16px; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th {{ background: #0d2a4f; color: #fff; padding: 6px; text-align: left; }}
    td {{ border-bottom: 1px solid #ccc; padding: 5px 6px; }}
    tr:nth-child(even) td {{ background: #f3f6fb; }}
    td.num {{ text-align: right; }}
    @media print {{
      body {{ margin: 0; }}
      th {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
    }}
  </style>
</head>
<body>
  <h1>Reporte de Lotes — CYS Ltda</h1>
  <div class="meta">Generado el {timezone.localdate():%d/%m/%Y} · {len(filas)} lote{plural}</div>
  <table>
    <thead><tr>{encabezados_html}</tr></thead>
    <tbody>{filas_html}</tbody>
  </table>
  <script>window.onload = function () {{ window.print(); }};</script>
</body>
</html>"""
    return HttpResponse(html)


@login_required
def lote_exportar(request, formato):
    if formato == 'excel':
        return _exportar_excel()
    if formato == 'pdf':
        return _exportar_pdf()
    if formato == 'imprimir':
        return _exportar_imprimir()
    raise Http404('Formato de exportación no válido.')