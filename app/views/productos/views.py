from decimal import Decimal, InvalidOperation
from io import BytesIO
from itertools import zip_longest
from django.core.exceptions import ValidationError
from app.validators import normalizar_nombre, validar_nombre_producto
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Prefetch, Sum
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape
from django.utils.http import url_has_allowed_host_and_scheme
from app.validators import normalizar_nombre, validar_nombre_producto, validar_nombre_presentacion
from app.forms import ProductoRegistroForm
from app.models import Categoria, Compra, PresentacionProducto, Producto

UMBRAL_STOCK_CRITICO = 5


# ===============================
# HELPERS
# ===============================
def _es_ajax(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'


def _destino_seguro(request, destino):
    """Evita open redirect: solo acepta rutas del mismo sitio."""
    if destino and url_has_allowed_host_and_scheme(
        destino,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return destino
    return reverse('lista_productos')


def _entero_positivo(valor):
    """Devuelve un entero >= 1, o None si no es válido."""
    try:
        n = int(str(valor).strip())
    except (TypeError, ValueError):
        return None
    return n if n >= 1 else None


def _decimal_positivo(valor):
    """Devuelve un Decimal > 0 con 2 decimales, o None si no es válido."""
    try:
        d = Decimal(str(valor).strip().replace(',', '.'))
    except (InvalidOperation, ValueError, AttributeError):
        return None
    if not d.is_finite() or d <= 0:
        return None
    return d.quantize(Decimal('0.01'))


# ===============================
# LISTA / VISTA PRINCIPAL
# ===============================
@login_required
def lista_productos(request):
    productos_qs = (
        Producto.objects.select_related('categoria')
        .prefetch_related('presentaciones__lotes', 'detalles')
        .all()
    )

    categorias = Categoria.objects.filter(subcategoria__isnull=True).prefetch_related(
        Prefetch(
            'productos',
            queryset=Producto.objects.prefetch_related('presentaciones__lotes').select_related('categoria'),
        ),
        Prefetch(
            'subcategorias',
            queryset=Categoria.objects.prefetch_related(
                Prefetch(
                    'productos',
                    queryset=Producto.objects.prefetch_related('presentaciones__lotes').select_related('categoria'),
                )
            ),
        ),
    )

    resumen_categorias = []
    for cat in Categoria.objects.filter(subcategoria__isnull=True):
        total = Producto.objects.filter(categoria=cat).count()
        total += Producto.objects.filter(categoria__subcategoria=cat).count()
        resumen_categorias.append({'pk': cat.pk, 'nombre': cat.nombre, 'total': total})

    context = {
        'productos': productos_qs,
        'categorias': categorias,
        'todas_cats': Categoria.objects.all(),
        'resumen_categorias': resumen_categorias,
        'form': ProductoRegistroForm(),
        'breadcrumb_items': [
            {'nombre': 'Inventario', 'url': None},
            {'nombre': 'Productos', 'url': None},
        ],
    }
    return render(request, 'productos/productos.html', context)


# ===============================
# CREAR PRODUCTO
# ===============================
@login_required
def crear_producto(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    ajax = _es_ajax(request)
    destino = _destino_seguro(request, request.POST.get('next') or request.GET.get('next'))

    categoria_pk = request.POST.get('categoria', '').strip()
    descripcion = request.POST.get('descripcion', '').strip()

    errores = {}
    try:
        nombre = validar_nombre_producto(request.POST.get('nombre', ''))
    except ValidationError as e:
        nombre = ''
        errores['nombre'] = e.messages

    categoria = None
    if not categoria_pk:
        errores['categoria'] = ['La categoría es obligatoria.']
    elif not categoria_pk.isdigit():
        errores['categoria'] = ['Categoría no válida.']
    else:
        categoria = Categoria.objects.filter(pk=int(categoria_pk), activo=True).first()
        if categoria is None:
            errores['categoria'] = ['La categoría no existe o está desactivada.']

    if errores:
        if ajax:
            return JsonResponse({'ok': False, 'errores': errores}, status=400)
        messages.error(request, 'Corrige los errores del formulario.')
        return redirect(destino)

    # La fecha de vencimiento se registra después, en Detalle de Producto.
    producto = Producto.objects.create(
        nombre=nombre,
        categoria=categoria,
        descripcion=descripcion,
    )

    messages.success(request, f'✅ Producto "{producto.nombre}" creado correctamente.')

    if ajax:
        return JsonResponse({'ok': True, 'pk': producto.pk, 'nombre': producto.nombre})
    return redirect(destino)


# ===============================
# DETALLE PRODUCTO
# ===============================
@login_required
def producto_detalle(request, pk):
    """
    Antes renderizaba productos.html sin contexto (página vacía).
    Ahora vuelve a la lista; ?producto=<pk> queda disponible para que el
    frontend resalte o abra ese producto.
    """
    get_object_or_404(Producto, pk=pk)
    return redirect(f"{reverse('lista_productos')}?producto={pk}")


# ===============================
# EDITAR PRODUCTO
# ===============================
@login_required
def producto_editar(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    if request.method != 'POST':
        return redirect('lista_productos')

    cambios = []
    avisos = []

    with transaction.atomic():
                # ── Datos básicos ──
        nombre_raw = request.POST.get('nombre')
        if nombre_raw is not None and normalizar_nombre(nombre_raw) != producto.nombre:
            try:
                producto.nombre = validar_nombre_producto(nombre_raw)
                cambios.append('nombre')
            except ValidationError as e:
                avisos.append(f'Nombre no actualizado: {e.messages[0]}')

        descripcion = request.POST.get('descripcion')
        if descripcion is not None:
            descripcion = descripcion.strip()
            if descripcion != (producto.descripcion or ''):
                producto.descripcion = descripcion
                cambios.append('descripción')

        categoria_pk = request.POST.get('categoria', '').strip()
        if categoria_pk.isdigit() and int(categoria_pk) != producto.categoria_id:
            nueva_cat = Categoria.objects.filter(pk=int(categoria_pk)).first()
            if nueva_cat:
                producto.categoria = nueva_cat
                cambios.append('categoría')
            else:
                avisos.append('La categoría elegida no existe.')

        if cambios:
            producto.save()

        # ── Presentaciones existentes ──
        # OJO: ya NO se copia el precio de venta al costo de los lotes.
        for key, valor in request.POST.items():
            if not key.startswith('pres_nombre_'):
                continue
            pres_id = key.replace('pres_nombre_', '')
            if not pres_id.isdigit():
                continue
            pres = PresentacionProducto.objects.filter(pk=int(pres_id), producto=producto).first()
            if pres is None:
                continue

            modificado = False

            if valor.strip() and normalizar_nombre(valor) != pres.nombre:
                try:
                    pres.nombre = validar_nombre_presentacion(valor)
                    modificado = True
                except ValidationError as e:
                    avisos.append(f'Nombre de presentación no actualizado: {e.messages[0]}')

            cant_raw = request.POST.get(f'pres_cantidad_{pres_id}', '').strip()
            if cant_raw:
                nueva_cant = _entero_positivo(cant_raw)
                if nueva_cant is None:
                    avisos.append(f'Cantidad inválida en "{pres.nombre}" (debe ser 1 o más).')
                elif nueva_cant != pres.cantidad:
                    pres.cantidad = nueva_cant
                    modificado = True

            precio_raw = request.POST.get(f'pres_precio_{pres_id}', '').strip()
            if precio_raw:
                nuevo_precio = _decimal_positivo(precio_raw)
                if nuevo_precio is None:
                    avisos.append(f'Precio inválido en "{pres.nombre}" (debe ser mayor a 0).')
                elif nuevo_precio != pres.precio_venta:
                    pres.precio_venta = nuevo_precio
                    modificado = True

            if modificado:
                pres.save()
                cambios.append(f'presentación "{pres.nombre}"')

        # ── Nuevas presentaciones ──
        nombres = request.POST.getlist('nueva_pres_nombre[]')
        cantidades = request.POST.getlist('nueva_pres_cantidad[]')
        precios = request.POST.getlist('nueva_pres_precio[]')

        for nombre_pres, cant_raw, precio_raw in zip_longest(nombres, cantidades, precios, fillvalue=''):
            if not nombre_pres.strip():
                continue
            try:
                nombre_pres = validar_nombre_presentacion(nombre_pres)
            except ValidationError as e:
                avisos.append(f'Presentación "{nombre_pres.strip()}" no creada: {e.messages[0]}')
                continue

            cantidad_pres = _entero_positivo(cant_raw)
            precio_pres = _decimal_positivo(precio_raw)
            if cantidad_pres is None or precio_pres is None:
                avisos.append(
                    f'Presentación "{nombre_pres}" no creada: cantidad y precio deben ser mayores a 0.'
                )
                continue

            if producto.presentaciones.filter(nombre__iexact=nombre_pres).exists():
                avisos.append(f'Presentación "{nombre_pres}" no creada: ya existe en este producto.')
                continue

            PresentacionProducto.objects.create(
                producto=producto,
                nombre=nombre_pres,
                cantidad=cantidad_pres,
                precio_venta=precio_pres,
            )
            cambios.append(f'nueva presentación "{nombre_pres}"')

    for aviso in avisos:
        messages.warning(request, f'⚠️ {aviso}')

    if cambios:
        messages.success(request, f"✅ Producto '{producto.nombre}' actualizado.")
    elif not avisos:
        messages.info(request, 'ℹ️ No se detectaron cambios en el producto.')

    return redirect('lista_productos')


# ===============================
# REGISTRO PRODUCTO
# ===============================
@login_required
def producto_registro(request):
    """Antes renderizaba productos.html sin contexto. Ahora guarda y vuelve a la lista."""
    if request.method == 'POST':
        form = ProductoRegistroForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ Producto registrado correctamente.')
        else:
            for campo, errs in form.errors.items():
                messages.error(request, f'{campo}: {", ".join(errs)}')
    return redirect('lista_productos')


# ===============================
# STOCK STATUS & NOTIFICACIONES
# ===============================
@login_required
def stock_status(request):
    """Retorna productos con stock crítico y compras pendientes."""
    productos = Producto.objects.filter(activo=True).annotate(total=Sum('lotes__stock_actual'))

    criticos = [
        {'nombre': p.nombre, 'total_stock': p.total or 0}
        for p in productos
        if (p.total or 0) <= UMBRAL_STOCK_CRITICO
    ]

    compras_pendientes_qs = (
        Compra.objects.filter(estado__in=['pendiente', 'confirmada'])
        .select_related('proveedor')
        .order_by('-fecha')[:20]
    )

    compras_pendientes = [
        {
            'id': c.codigo_compra,
            'proveedor': c.proveedor.nombre_empresa if c.proveedor else 'Proveedor',
            'estado': c.get_estado_display(),
            'estado_raw': c.estado,
            'valor': float(c.valor),
            'saldo': float(c.saldo),
            'fecha': timezone.localtime(c.fecha).strftime('%d/%m %H:%M') if c.fecha else '',
            'url': reverse('detalle_compra', args=[c.codigo_compra]),
        }
        for c in compras_pendientes_qs
    ]

    return JsonResponse({
        'criticos': criticos,
        'compras_pendientes': compras_pendientes,
        'total_compras_pendientes': len(compras_pendientes),
        'total_alertas': len(criticos) + len(compras_pendientes),
    })


# ===============================
# BUSCAR PRODUCTO
# ===============================
@login_required
def buscar_producto(request):
    q = request.GET.get('q', '').strip()
    modo = request.GET.get('modo', '')

    if not q:
        if modo == 'sugerencias':
            return JsonResponse({'resultados': []})
        return JsonResponse({'encontrado': False, 'mensaje': 'Escribe un nombre.'})

    if modo == 'sugerencias':
        productos = Producto.objects.filter(nombre__icontains=q).select_related('categoria')[:8]
        resultados = [{
            'pk': p.pk,
            'nombre': p.nombre,
            'categoria': p.categoria.nombre if p.categoria else '—',
        } for p in productos]
        return JsonResponse({'resultados': resultados})

    producto = Producto.objects.filter(nombre__icontains=q).order_by('nombre').first()

    if not producto:
        return JsonResponse({'encontrado': False, 'mensaje': f'No se encontró "{q}".'})

    stock_total = producto.presentaciones.aggregate(
        total=Sum('lotes__stock_actual')
    )['total'] or 0

    presentaciones = []
    for pres in producto.presentaciones.all():
        stock_pres = pres.lotes.aggregate(total=Sum('stock_actual'))['total'] or 0
        presentaciones.append({
            'id': pres.pk,
            'nombre': pres.nombre,
            'cantidad': pres.cantidad,
            'precio': str(pres.precio_venta),
            'stock_actual': stock_pres,
        })

    return JsonResponse({
        'encontrado': True,
        'producto': {
            'pk': producto.pk,
            'nombre': producto.nombre,
            'categoria': producto.categoria.nombre if producto.categoria else '—',
            'stock_total': stock_total,
            'descripcion': producto.descripcion or '',
            'presentaciones': presentaciones,
        },
    })


# ===============================
# ACTIVAR / DESACTIVAR
# ===============================
@login_required
def producto_toggle_activo(request, pk):
    # TODO: pasar a @require_POST cuando el template/JS envíe POST con CSRF.
    producto = get_object_or_404(Producto, pk=pk)
    producto.activo = not producto.activo
    producto.save(update_fields=['activo'])
    estado = 'activado' if producto.activo else 'desactivado'
    if _es_ajax(request):
        return JsonResponse({'ok': True, 'nombre': producto.nombre, 'activo': producto.activo})
    messages.success(request, f'Producto "{producto.nombre}" {estado}.')
    return redirect('lista_productos')


# ===============================
# EXPORTAR PRODUCTOS (PDF / EXCEL / IMPRIMIR)
# ===============================
ENCABEZADOS_EXPORT_PROD = [
    '#', 'Producto', 'Categoría', 'Código', 'Presentaciones', 'Stock total', 'Estado',
]


def _presentaciones_texto(producto):
    """Lista de textos, uno por presentación: 'Caja (x24) · $50,000'."""
    textos = []
    for pres in producto.presentaciones.all():
        precio = pres.precio_venta or 0
        textos.append(f'{pres.nombre} (x{pres.cantidad}) · ${precio:,.0f}')
    return textos


def _filas_productos():
    """Devuelve las filas (listas) de todos los productos para exportar."""
    productos = (
        Producto.objects
        .select_related('categoria')
        .prefetch_related('presentaciones__lotes', 'detalles')
        .order_by('nombre')
    )
    filas = []
    for i, p in enumerate(productos, start=1):
        detalles = list(p.detalles.all())
        codigo = detalles[0].codigo_barras if detalles and detalles[0].codigo_barras else ''
        filas.append([
            i,
            p.nombre,
            p.categoria.nombre if p.categoria else '—',
            codigo,
            _presentaciones_texto(p),
            p.stock_total or 0,
            'Activo' if p.activo else 'Inactivo',
        ])
    return filas


def _nombre_archivo_prod(extension):
    return f'productos_{timezone.localdate():%Y%m%d}.{extension}'


def _exportar_excel_productos():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Productos'

    ws.append(ENCABEZADOS_EXPORT_PROD)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0D2A4F')
        celda.alignment = Alignment(horizontal='center', vertical='center')

    for fila in _filas_productos():
        fila[4] = '\n'.join(fila[4])
        ws.append(fila)

    # Ajuste de texto en la columna de presentaciones y anchos de columna
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for celda in row:
            celda.alignment = Alignment(wrap_text=True, vertical='top')

    for i, titulo in enumerate(ENCABEZADOS_EXPORT_PROD, start=1):
        largos = [len(str(titulo))]
        for c in ws[get_column_letter(i)][1:]:
            for linea in str(c.value or '').split('\n'):
                largos.append(len(linea))
        ws.column_dimensions[get_column_letter(i)].width = min(max(largos) + 3, 50)
    ws.freeze_panes = 'A2'

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo_prod("xlsx")}"'
    return response


def _exportar_pdf_productos():
    from xml.sax.saxutils import escape as xml_escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm,
        title='Reporte de Productos',
    )
    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle('celda', parent=estilos['Normal'], fontSize=8, leading=10)

    datos = [ENCABEZADOS_EXPORT_PROD]
    for f in _filas_productos():
        presentaciones = '<br/>'.join(xml_escape(t) for t in f[4]) or '—'
        datos.append([
            f[0],
            Paragraph(xml_escape(f[1]), estilo_celda),
            Paragraph(xml_escape(f[2]), estilo_celda),
            f[3],
            Paragraph(presentaciones, estilo_celda),
            f[5],
            f[6],
        ])

    tabla = Table(
        datos, repeatRows=1,
        colWidths=[1.0 * cm, 6.0 * cm, 4.0 * cm, 3.5 * cm, 8.0 * cm, 2.2 * cm, 2.0 * cm],
    )
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D2A4F')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (5, 1), (5, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF3FA')]),
    ]))

    elementos = [
        Paragraph('Reporte de Productos — CYS Ltda', estilos['Title']),
        Paragraph(f'Generado el {timezone.localdate():%d/%m/%Y}', estilos['Normal']),
        Spacer(1, 0.4 * cm),
        tabla,
    ]
    doc.build(elementos)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo_prod("pdf")}"'
    return response


def _exportar_imprimir_productos():
    filas = _filas_productos()

    encabezados_html = ''.join(f'<th>{escape(e)}</th>' for e in ENCABEZADOS_EXPORT_PROD)

    filas_html = ''
    for f in filas:
        presentaciones = '<br>'.join(escape(t) for t in f[4]) or '—'
        filas_html += (
            '<tr>'
            f'<td>{f[0]}</td>'
            f'<td>{escape(f[1])}</td>'
            f'<td>{escape(f[2])}</td>'
            f'<td>{escape(f[3])}</td>'
            f'<td>{presentaciones}</td>'
            f'<td class="num">{f[5]}</td>'
            f'<td>{escape(f[6])}</td>'
            '</tr>'
        )
    if not filas_html:
        filas_html = '<tr><td colspan="7">No hay productos registrados.</td></tr>'

    plural = '' if len(filas) == 1 else 's'
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <title>Reporte de Productos — CYS Ltda</title>
  <style>
    body {{ font-family: Arial, sans-serif; font-size: 12px; color: #111; margin: 24px; }}
    h1 {{ font-size: 20px; margin: 0 0 4px; }}
    .meta {{ color: #555; margin-bottom: 16px; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th {{ background: #0d2a4f; color: #fff; padding: 6px; text-align: left; }}
    td {{ border-bottom: 1px solid #ccc; padding: 5px 6px; vertical-align: top; }}
    tr:nth-child(even) td {{ background: #f3f6fb; }}
    td.num {{ text-align: right; }}
    @media print {{
      body {{ margin: 0; }}
      th {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
    }}
  </style>
</head>
<body>
  <h1>Reporte de Productos — CYS Ltda</h1>
  <div class="meta">Generado el {timezone.localdate():%d/%m/%Y} · {len(filas)} producto{plural}</div>
  <table>
    <thead><tr>{encabezados_html}</tr></thead>
    <tbody>{filas_html}</tbody>
  </table>
  <script>window.onload = function () {{ window.print(); }};</script>
</body>
</html>"""
    return HttpResponse(html)


@login_required
def producto_exportar(request, formato):
    if formato == 'excel':
        return _exportar_excel_productos()
    if formato == 'pdf':
        return _exportar_pdf_productos()
    if formato == 'imprimir':
        return _exportar_imprimir_productos()
    raise Http404('Formato de exportación no válido.')