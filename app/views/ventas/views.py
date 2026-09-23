import json
from decimal import Decimal, InvalidOperation
from functools import wraps
from datetime import datetime, timedelta

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse
from django.db import transaction
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import csrf_exempt

from app.models import (
    Venta, DetalleVenta, Caja, Devolucion, DetalleDevolucion,
    Producto, Categoria, PresentacionProducto, Usuario
)

BILLETES_DENOM = [100000, 50000, 20000, 10000, 5000, 2000, 1000]
MONEDAS_DENOM  = [500, 200, 100, 50]


# ════════════════════════════════════════
# DECORADOR
# ════════════════════════════════════════

def session_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated and not request.session.get('usuario_id'):
            return redirect('login')
        return view_func(request, *args, **kwargs)
    return wrapper


# ════════════════════════════════════════
# VENTAS
# ════════════════════════════════════════

def obtener_caja_abierta(hoy=None):
    caja_ultima = Caja.objects.order_by('-codigo_caja').first()
    if caja_ultima and '[CERRADO]' not in (caja_ultima.observacion or ''):
        return caja_ultima
    return None


@session_required
def ventas_lista(request):
    ventas = Venta.objects.all().order_by('-fecha')
    categorias = Categoria.objects.all()
    hoy = timezone.localdate()
    ventas_hoy = Venta.objects.filter(fecha__date=hoy)
    total_dia = int(sum(v.total_venta for v in ventas_hoy))

    caja_abierta = obtener_caja_abierta(hoy)
    ultimo_cierre = Caja.objects.order_by('-fecha_hora').first()

    return render(request, 'ventas/ventas.html', {
        'ventas': ventas,
        'categorias': categorias,
        'total_dia': total_dia,
        'hoy': hoy,
        'caja_abierta': caja_abierta,
        'ultimo_cierre': ultimo_cierre,
        'billetes_denom': BILLETES_DENOM,
        'monedas_denom': MONEDAS_DENOM,
    })


# Alias para compatibilidad con urls.py y aside.html
ventas = ventas_lista


@session_required
def nueva_venta(request):
    if request.method != 'POST':
        return redirect('ventas')

    vendedor = obtener_usuario_actual(request)

    caja_actual = obtener_caja_abierta()
    if not caja_actual:
        caja_actual = Caja.objects.create(
            monto_base=Decimal('0'),
            total_efectivo=Decimal('0'),
            total_transferencias=Decimal('0'),
            total_retirado=Decimal('0'),
            usuario=vendedor,
            observacion='Turno iniciado automáticamente al registrar venta',
        )

    producto_ids = request.POST.getlist('producto_id[]')
    presentacion_ids = request.POST.getlist('presentacion_id[]')
    cantidades = request.POST.getlist('cantidad[]')
    precios = request.POST.getlist('precio[]')

    def to_decimal(key, default='0'):
        try:
            return Decimal(request.POST.get(key, default) or default)
        except (InvalidOperation, TypeError):
            return Decimal('0')

    descuento_pct = to_decimal('descuento_porcentaje')
    pago_efectivo = to_decimal('pago_efectivo')
    pago_tarjeta = to_decimal('pago_tarjeta')
    pago_transferencia = to_decimal('pago_transferencia')
    pago_nequi = to_decimal('pago_nequi')
    pago_daviplata = to_decimal('pago_daviplata')

    if not producto_ids:
        messages.error(request, "El carrito está vacío.")
        return redirect('ventas')

    items_validados = []
    subtotal_venta = Decimal('0')

    for i, prod_id in enumerate(producto_ids):
        try:
            cantidad = int(cantidades[i])
            precio = Decimal(precios[i])
            if cantidad <= 0 or precio < 0:
                raise ValueError
        except (ValueError, TypeError, InvalidOperation, IndexError):
            messages.error(request, f"Datos inválidos en el ítem {i+1}.")
            return redirect('ventas')

        try:
            producto = Producto.objects.get(pk=prod_id)
        except Producto.DoesNotExist:
            messages.error(request, f"Producto {i+1} no encontrado.")
            return redirect('ventas')

        pres_id = presentacion_ids[i] if i < len(presentacion_ids) else ''
        presentacion = None

        if pres_id and pres_id != 'null':
            try:
                presentacion = PresentacionProducto.objects.get(pk=pres_id)
            except PresentacionProducto.DoesNotExist:
                messages.error(request, f"Presentación inválida para {producto.nombre}.")
                return redirect('ventas')

        items_validados.append({
            'producto': producto,
            'presentacion': presentacion,
            'cantidad': cantidad,
            'precio': precio,
        })
        subtotal_venta += precio * cantidad

    monto_descuento = (subtotal_venta * descuento_pct) / Decimal('100')
    total_final = subtotal_venta - monto_descuento
    total_pagado = pago_efectivo + pago_tarjeta + pago_transferencia + pago_nequi + pago_daviplata

    if total_pagado > Decimal('0') and total_pagado < total_final:
        messages.error(request, f"El total pagado (${total_pagado:,.0f}) no cubre el total (${total_final:,.0f}).".replace(',', '.'))
        return redirect('ventas')

    metodo_pago = 'efectivo'
    if pago_tarjeta > 0:
        metodo_pago = 'tarjeta'
    elif pago_transferencia > 0:
        metodo_pago = 'transferencia'
    elif pago_nequi > 0:
        metodo_pago = 'nequi'
    elif pago_daviplata > 0:
        metodo_pago = 'daviplata'

    venta = Venta.objects.create(
        total_venta=total_final,
        metodo_pago=metodo_pago,
        usuario=vendedor,
        caja=caja_actual,
    )

    for item in items_validados:
        DetalleVenta.objects.create(
            venta=venta,
            producto=item['producto'],
            cantidad=item['cantidad'],
            precio_unitario=item['precio'],
            subtotal=item['precio'] * item['cantidad'],
            total=item['precio'] * item['cantidad'],
        )

    messages.success(request, f"Venta registrada — Total: ${total_final:,.0f}".replace(',', '.'))
    return redirect('ventas')


@session_required
def eliminar_venta(request, pk):
    venta = get_object_or_404(Venta, pk=pk)
    if request.method == 'POST':
        venta.delete()
        messages.success(request, "Venta eliminada.")
    return redirect('ventas')


def producto_stock_json(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    presentaciones = []
    if hasattr(producto, 'presentaciones'):
        presentaciones = [
            {'id': p.id, 'nombre': p.nombre, 'cantidad': getattr(p, 'cantidad', 0), 'precio': float(getattr(p, 'precio_venta', 0))}
            for p in producto.presentaciones.all()
        ]
    return JsonResponse({
        'stock': getattr(producto, 'cantidad_disponible', 0),
        'precio': float(getattr(producto, 'precio_unitario', 0)),
        'unidad': getattr(producto, 'unidad', 'Unid'),
        'presentaciones': presentaciones,
    })


# ════════════════════════════════════════
# VENTAS DEL DÍA
# ════════════════════════════════════════

@session_required
def ventas_dia(request):
    hoy = timezone.localdate()
    fecha_inicio_str = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_str = request.GET.get('fecha_fin', '').strip()

    fecha_inicio = hoy
    fecha_fin = hoy

    if fecha_inicio_str:
        try:
            fecha_inicio = datetime.strptime(fecha_inicio_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    if fecha_fin_str:
        try:
            fecha_fin = datetime.strptime(fecha_fin_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    if fecha_inicio_str or fecha_fin_str:
        ventas_qs = Venta.objects.filter(fecha__date__range=[fecha_inicio, fecha_fin]).order_by('-fecha')
    else:
        ventas_qs = Venta.objects.filter(fecha__date=hoy).order_by('-fecha')

    ventas_list = list(ventas_qs.select_related('vendedor').prefetch_related('detalles__producto'))
    total_dia = float(sum(v.total_venta for v in ventas_list))
    total_productos = sum(det.cantidad for v in ventas_list for det in v.detalles.all())

    # Tope de ventas por periodo (Día, Semana, Mes, Año)
    inicio_semana = hoy - timedelta(days=hoy.weekday())
    inicio_mes = hoy.replace(day=1)
    inicio_ano = hoy.replace(month=1, day=1)

    v_dia_agg = Venta.objects.filter(fecha__date=hoy).aggregate(total=Sum('total_venta'), cant=Count('pk'))
    total_dia_val = float(v_dia_agg['total'] or 0)
    cant_dia = v_dia_agg['cant'] or 0

    v_sem_agg = Venta.objects.filter(fecha__date__range=[inicio_semana, hoy]).aggregate(total=Sum('total_venta'), cant=Count('pk'))
    total_semana_val = float(v_sem_agg['total'] or 0)
    cant_semana = v_sem_agg['cant'] or 0

    v_mes_agg = Venta.objects.filter(fecha__date__range=[inicio_mes, hoy]).aggregate(total=Sum('total_venta'), cant=Count('pk'))
    total_mes_val = float(v_mes_agg['total'] or 0)
    cant_mes = v_mes_agg['cant'] or 0

    v_ano_agg = Venta.objects.filter(fecha__date__range=[inicio_ano, hoy]).aggregate(total=Sum('total_venta'), cant=Count('pk'))
    total_ano_val = float(v_ano_agg['total'] or 0)
    cant_ano = v_ano_agg['cant'] or 0

    meses_es = ['', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
    nombre_mes = meses_es[hoy.month]

    base_tope = max(total_ano_val, 1)

    topes_ventas = [
        {
            'key': 'dia',
            'label': 'Ventas Hoy',
            'sublabel': hoy.strftime('%d/%m/%Y'),
            'icon': 'bi-sun-fill',
            'color': '#2ed573',
            'badge_bg': 'rgba(46,213,115,0.15)',
            'monto': total_dia_val,
            'cantidad': cant_dia,
            'porcentaje': round((total_dia_val / base_tope * 100), 1) if base_tope > 0 else 0,
        },
        {
            'key': 'semana',
            'label': 'Esta Semana',
            'sublabel': f"{inicio_semana.strftime('%d/%m')} - {hoy.strftime('%d/%m')}",
            'icon': 'bi-calendar-week-fill',
            'color': '#4DA8DA',
            'badge_bg': 'rgba(77,168,218,0.15)',
            'monto': total_semana_val,
            'cantidad': cant_semana,
            'porcentaje': round((total_semana_val / base_tope * 100), 1) if base_tope > 0 else 0,
        },
        {
            'key': 'mes',
            'label': 'Este Mes',
            'sublabel': f"{nombre_mes} {hoy.year}",
            'icon': 'bi-calendar-month-fill',
            'color': '#a55eea',
            'badge_bg': 'rgba(165,94,234,0.15)',
            'monto': total_mes_val,
            'cantidad': cant_mes,
            'porcentaje': round((total_mes_val / base_tope * 100), 1) if base_tope > 0 else 0,
        },
        {
            'key': 'ano',
            'label': 'Este Año',
            'sublabel': f"Año {hoy.year}",
            'icon': 'bi-trophy-fill',
            'color': '#ffab00',
            'badge_bg': 'rgba(255,171,0,0.15)',
            'monto': total_ano_val,
            'cantidad': cant_ano,
            'porcentaje': 100.0 if total_ano_val > 0 else 0,
        },
    ]

    # Ranking de productos más vendidos
    prods_map = {}
    for v in ventas_list:
        for det in v.detalles.all():
            p_id = det.producto.pk
            p_nombre = det.producto.nombre
            if p_id not in prods_map:
                prods_map[p_id] = {
                    'nombre': p_nombre,
                    'cantidad': 0,
                    'recaudo': 0.0,
                }
            prods_map[p_id]['cantidad'] += det.cantidad
            prods_map[p_id]['recaudo'] += float(det.subtotal)

    productos_top = sorted(prods_map.values(), key=lambda x: x['cantidad'], reverse=True)[:8]

    max_cant_top = productos_top[0]['cantidad'] if productos_top else 1
    for i, p in enumerate(productos_top, start=1):
        p['posicion'] = i
        p['porcentaje'] = round((p['cantidad'] / total_productos * 100), 1) if total_productos > 0 else 0
        p['relativo'] = round((p['cantidad'] / max_cant_top * 100), 1) if max_cant_top > 0 else 0

    caja_abierta = obtener_caja_abierta(hoy)
    ultimo_cierre = Caja.objects.order_by('-fecha_hora').first()

    return render(request, 'ventas/ventas_dia.html', {
        'ventas': ventas_qs,
        'total_dia': total_dia,
        'total_productos': total_productos,
        'hoy': hoy,
        'fecha_inicio': fecha_inicio_str,
        'fecha_fin': fecha_fin_str,
        'topes_ventas': topes_ventas,
        'productos_top': productos_top,
        'caja_abierta': caja_abierta,
        'ultimo_cierre': ultimo_cierre,
    })


# ════════════════════════════════════════
# CAJA — APERTURA Y CIERRE
# ════════════════════════════════════════

@session_required
def caja(request):
    cajas_asc = list(Caja.objects.all().order_by('fecha_hora'))
    turnos = []
    turnos_dict_for_js = {}

    for i, item in enumerate(cajas_asc):
        inicio_turno = item.fecha_hora
        if i < len(cajas_asc) - 1:
            fin_turno = cajas_asc[i+1].fecha_hora
        else:
            fin_turno = timezone.now()

        # Obtener ventas asociadas al turno (por FK o por intervalo de fecha)
        ventas_turno = Venta.objects.filter(
            Q(caja=item) | (Q(fecha__gte=inicio_turno) & Q(fecha__lte=fin_turno))
        ).distinct()

        total_ventas_turno = Decimal('0')
        cant_ventas = ventas_turno.count()

        pago_efectivo = Decimal('0')
        pago_tarjeta = Decimal('0')
        pago_transferencia = Decimal('0')
        pago_nequi = Decimal('0')
        pago_daviplata = Decimal('0')

        for v in ventas_turno:
            monto_v = Decimal(str(v.total_con_descuento if (v.total_con_descuento and v.total_con_descuento > 0) else (v.total_venta or 0)))
            total_ventas_turno += monto_v

            pe = Decimal(str(v.pago_efectivo or 0))
            pt = Decimal(str(v.pago_tarjeta or 0))
            ptr = Decimal(str(v.pago_transferencia or 0))
            pn = Decimal(str(v.pago_nequi or 0))
            pd = Decimal(str(v.pago_daviplata or 0))

            if pe > 0 or pt > 0 or ptr > 0 or pn > 0 or pd > 0:
                pago_efectivo += pe
                pago_tarjeta += pt
                pago_transferencia += ptr
                pago_nequi += pn
                pago_daviplata += pd
            else:
                mp = (v.metodo_pago or 'efectivo').lower()
                if 'efectivo' in mp:
                    pago_efectivo += monto_v
                elif 'tarjeta' in mp:
                    pago_tarjeta += monto_v
                elif 'nequi' in mp:
                    pago_nequi += monto_v
                elif 'daviplata' in mp:
                    pago_daviplata += monto_v
                else:
                    pago_transferencia += monto_v

        # Productos top en este turno
        detalles_turno = DetalleVenta.objects.filter(venta__in=ventas_turno).select_related('producto', 'presentacion')
        productos_dict = {}
        for d in detalles_turno:
            p_nombre = d.producto.nombre if d.producto else (d.presentacion.nombre if d.presentacion else 'Producto Sin Nombre')
            cant = d.cantidad or 0
            subt = float((d.cantidad or 0) * (d.precio_unitario or 0)) if hasattr(d, 'precio_unitario') else float(d.subtotal() if callable(getattr(d, 'subtotal', None)) else 0)

            if p_nombre not in productos_dict:
                productos_dict[p_nombre] = {'nombre': p_nombre, 'cantidad': 0, 'total': 0.0}
            productos_dict[p_nombre]['cantidad'] += cant
            productos_dict[p_nombre]['total'] += subt

        total_cant_prods = sum(p['cantidad'] for p in productos_dict.values())
        productos_top_turno = sorted(productos_dict.values(), key=lambda x: x['cantidad'], reverse=True)[:5]

        max_cant_prod = productos_top_turno[0]['cantidad'] if productos_top_turno else 1
        for pos, p in enumerate(productos_top_turno, start=1):
            p['posicion'] = pos
            p['porcentaje_total'] = round((p['cantidad'] / total_cant_prods * 100), 1) if total_cant_prods > 0 else 0
            p['porcentaje_relativo'] = round((p['cantidad'] / max_cant_prod * 100), 1) if max_cant_prod > 0 else 0

        # Estado del turno
        es_cerrado = '[CERRADO]' in (item.observacion or '')
        es_ultimo = (i == len(cajas_asc) - 1)
        estado_label = 'CERRADO' if es_cerrado else ('ACTIVO' if es_ultimo else 'FINALIZADO')

        # Operador
        if item.usuario:
            operador_nombre = getattr(item.usuario, 'nombre', None) or item.usuario.get_full_name() or item.usuario.username or str(item.usuario.documento_usuario or '')
        else:
            operador_nombre = 'Administrador / Sistema'

        monto_base_val = float(item.monto_base or 0)
        total_ventas_val = float(total_ventas_turno)
        monto_total_acumulado = monto_base_val + total_ventas_val

        pct_base = round((monto_base_val / monto_total_acumulado * 100), 1) if monto_total_acumulado > 0 else (100.0 if monto_base_val > 0 else 0.0)
        pct_generado = round((total_ventas_val / monto_total_acumulado * 100), 1) if monto_total_acumulado > 0 else 0.0

        pct_efectivo = round((float(pago_efectivo) / total_ventas_val * 100), 1) if total_ventas_val > 0 else 0.0
        pct_nequi = round((float(pago_nequi) / total_ventas_val * 100), 1) if total_ventas_val > 0 else 0.0
        pct_daviplata = round((float(pago_daviplata) / total_ventas_val * 100), 1) if total_ventas_val > 0 else 0.0
        pct_tarjeta = round((float(pago_tarjeta) / total_ventas_val * 100), 1) if total_ventas_val > 0 else 0.0
        pct_transferencia = round((float(pago_transferencia) / total_ventas_val * 100), 1) if total_ventas_val > 0 else 0.0

        turno_dict = {
            'id': item.codigo_caja,
            'caja_obj': item,
            'fecha_inicio': inicio_turno,
            'fecha_fin': fin_turno if es_cerrado or not es_ultimo else None,
            'monto_base': item.monto_base or Decimal('0'),
            'total_ventas': total_ventas_turno,
            'monto_total_acumulado': Decimal(str(monto_total_acumulado)),
            'pct_base': pct_base,
            'pct_generado': pct_generado,
            'cant_ventas': cant_ventas,
            'total_efectivo_arqueo': item.total_efectivo or Decimal('0'),
            'total_transferencias_arqueo': item.total_transferencias or Decimal('0'),
            'total_retirado': item.total_retirado or Decimal('0'),
            'pago_efectivo': pago_efectivo,
            'pago_tarjeta': pago_tarjeta,
            'pago_transferencia': pago_transferencia,
            'pago_nequi': pago_nequi,
            'pago_daviplata': pago_daviplata,
            'pct_efectivo': pct_efectivo,
            'pct_nequi': pct_nequi,
            'pct_daviplata': pct_daviplata,
            'pct_tarjeta': pct_tarjeta,
            'pct_transferencia': pct_transferencia,
            'productos_top': productos_top_turno,
            'denominaciones': item.denominaciones or {},
            'observacion': (item.observacion or '').replace('[CERRADO]', '').strip(),
            'operador': operador_nombre,
            'estado': estado_label
        }

        turnos.append(turno_dict)

        # Copia JSON para el modal interactivo
        turnos_dict_for_js[str(item.codigo_caja)] = {
            'id': item.codigo_caja,
            'fecha_inicio': inicio_turno.strftime('%d/%m/%Y %I:%M %p') if inicio_turno else '',
            'fecha_fin': fin_turno.strftime('%d/%m/%Y %I:%M %p') if (es_cerrado or not es_ultimo) and fin_turno else ('En Curso' if es_ultimo and not es_cerrado else ''),
            'monto_base': float(item.monto_base or 0),
            'total_ventas': float(total_ventas_turno),
            'monto_total_acumulado': float(monto_total_acumulado),
            'pct_base': pct_base,
            'pct_generado': pct_generado,
            'cant_ventas': cant_ventas,
            'total_efectivo_arqueo': float(item.total_efectivo or 0),
            'total_transferencias_arqueo': float(item.total_transferencias or 0),
            'total_retirado': float(item.total_retirado or 0),
            'pago_efectivo': float(pago_efectivo),
            'pago_tarjeta': float(pago_tarjeta),
            'pago_transferencia': float(pago_transferencia),
            'pago_nequi': float(pago_nequi),
            'pago_daviplata': float(pago_daviplata),
            'pct_efectivo': pct_efectivo,
            'pct_nequi': pct_nequi,
            'pct_daviplata': pct_daviplata,
            'pct_tarjeta': pct_tarjeta,
            'pct_transferencia': pct_transferencia,
            'productos_top': productos_top_turno,
            'denominaciones': item.denominaciones or {},
            'observacion': (item.observacion or '').replace('[CERRADO]', '').strip(),
            'operador': operador_nombre,
            'estado': estado_label
        }

    turnos.reverse()

    total_turnos = len(turnos)
    turno_activo = next((t for t in turnos if t['estado'] == 'ACTIVO'), None)
    caja_abierta = obtener_caja_abierta()
    ultimo_cierre = Caja.objects.order_by('-fecha_hora').first()
    ventas_totales_historicas = sum(t['total_ventas'] for t in turnos)
    promedio_base = (sum(t['monto_base'] for t in turnos) / total_turnos) if total_turnos > 0 else 0

    return render(request, 'ventas/caja.html', {
        'turnos': turnos,
        'turnos_js': json.dumps(turnos_dict_for_js),
        'total_turnos': total_turnos,
        'turno_activo': turno_activo,
        'caja_abierta': caja_abierta,
        'ultimo_cierre': ultimo_cierre,
        'ventas_totales_historicas': ventas_totales_historicas,
        'promedio_base': promedio_base,
        'billetes_denom': BILLETES_DENOM,
        'monedas_denom': MONEDAS_DENOM,
    })


def obtener_usuario_actual(request):
    vendedor = None
    if hasattr(request, 'user') and request.user and request.user.is_authenticated:
        try:
            vendedor = Usuario.objects.get(pk=request.user.pk)
        except (AttributeError, Usuario.DoesNotExist):
            if hasattr(request.user, '_wrapped') and isinstance(request.user._wrapped, Usuario):
                vendedor = request.user._wrapped

    if not vendedor and request.session.get('usuario_id'):
        vendedor = Usuario.objects.filter(pk=request.session.get('usuario_id')).first()

    if not vendedor or not isinstance(vendedor, Usuario):
        vendedor = Usuario.objects.first()

    return vendedor


@csrf_exempt
@session_required
def apertura_caja(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'JSON inválido: {str(e)}'}, status=400)

    try:
        monto_base = float(data.get('monto_base', 0) or data.get('monto_contado', 0) or 0)
        if monto_base < 0:
            monto_base = 0.0
    except (TypeError, ValueError):
        monto_base = 0.0

    vendedor = obtener_usuario_actual(request)
    observacion_clean = (data.get('observacion', '') or '').replace('[CERRADO]', '').strip()

    try:
        caja_existente = obtener_caja_abierta()
        if caja_existente:
            caja_existente.monto_base = Decimal(str(monto_base))
            caja_existente.observacion = observacion_clean
            caja_existente.denominaciones = data.get('denominaciones', {})
            if vendedor:
                caja_existente.usuario = vendedor
            caja_existente.save()
            codigo_caja = caja_existente.codigo_caja
        else:
            nueva_caja = Caja.objects.create(
                fecha_hora=timezone.now(),
                monto_base=Decimal(str(monto_base)),
                total_efectivo=Decimal('0'),
                total_transferencias=Decimal('0'),
                total_retirado=Decimal('0'),
                usuario=vendedor,
                observacion=observacion_clean,
                denominaciones=data.get('denominaciones', {}),
            )
            codigo_caja = nueva_caja.codigo_caja
        return JsonResponse({'ok': True, 'monto_base': monto_base, 'codigo_caja': codigo_caja})
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error al guardar apertura: {str(e)}'}, status=500)


@csrf_exempt
@session_required
def cierre_caja(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'JSON inválido: {str(e)}'}, status=400)

    caja_reg = obtener_caja_abierta()
    if not caja_reg:
        caja_reg = Caja.objects.order_by('-fecha_hora').first()

    if not caja_reg:
        return JsonResponse({'ok': False, 'error': 'No hay registro de caja disponible para cerrar.'}, status=400)

    try:
        total_contado = float(data.get('total_contado', 0))
        total_retirado = float(data.get('total_retirado', 0))
    except (TypeError, ValueError):
        total_contado = 0.0
        total_retirado = 0.0

    try:
        caja_reg.total_efectivo = total_contado
        caja_reg.total_retirado = total_retirado
        obs = (caja_reg.observacion or '').strip()
        if '[CERRADO]' not in obs:
            caja_reg.observacion = (obs + ' [CERRADO]').strip()

        if 'denominaciones' in data:
            caja_reg.denominaciones = data['denominaciones']
        caja_reg.save()

        return JsonResponse({'ok': True})
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error al guardar cierre: {str(e)}'}, status=500)


@require_POST
@session_required
def registrar_conteo(request):
    return apertura_caja(request)


# ════════════════════════════════════════
# DEVOLUCIONES
# ════════════════════════════════════════

@session_required
def lista_devoluciones(request):
    devoluciones = Devolucion.objects.all().order_by('-fecha')
    return render(request, 'ventas/devoluciones.html', {'devoluciones': devoluciones})


@session_required
def buscar_venta_devolucion(request):
    q = request.GET.get('q', '').strip()
    if not q:
        return JsonResponse({'ventas': []})

    ventas_qs = Venta.objects.all().order_by('-fecha')
    if q.isdigit():
        ventas_qs = Venta.objects.filter(pk=int(q))

    return JsonResponse({'ventas': [
        {'id': v.pk, 'cliente': str(getattr(v, 'usuario', 'Cliente')),
         'fecha': v.fecha.strftime('%d/%m/%Y %H:%M'), 'total': float(v.total_venta)}
        for v in ventas_qs[:10]
    ]})


@session_required
def detalle_venta_devolucion(request, venta_id):
    venta = get_object_or_404(Venta, pk=venta_id)
    return JsonResponse({
        'venta_id': venta.pk,
        'cliente': str(getattr(venta, 'usuario', 'Cliente')),
        'fecha': venta.fecha.strftime('%d/%m/%Y %H:%M'),
        'total': float(venta.total_venta),
        'descuento': 0,
        'detalles': [
            {
                'detalle_id': d.pk,
                'producto_id': d.producto.pk,
                'producto': d.producto.nombre,
                'cantidad': d.cantidad,
                'precio': float(d.precio_unitario),
                'subtotal': float(d.subtotal),
            }
            for d in getattr(venta, 'detalles', [])
        ],
    })


@session_required
@transaction.atomic
def registrar_devolucion(request):
    if request.method != 'POST':
        return redirect('lista_devoluciones')

    venta_id = request.POST.get('venta_id')
    tiene_comprobante = request.POST.get('tiene_comprobante') == '1'
    motivo = request.POST.get('motivo', 'otro')
    observaciones = request.POST.get('observaciones', '')
    producto_ids = request.POST.getlist('producto_id[]')
    cantidades = request.POST.getlist('cantidad[]')
    precios = request.POST.getlist('precio[]')

    if not tiene_comprobante:
        messages.warning(request, 'No se puede registrar la devolución sin comprobante de compra.')
        return redirect('lista_devoluciones')

    venta = get_object_or_404(Venta, pk=venta_id)
    total_devuelto = sum(int(cantidades[i]) * float(precios[i]) for i in range(len(producto_ids)))

    devolucion = Devolucion.objects.create(
        venta=venta,
        motivo=motivo,
        tipo_devolucion='devolucion',
        observaciones=observaciones,
        presenta_comprobante=tiene_comprobante,
        total_devuelto=total_devuelto,
        estado='completado',
        cantidad_cambio=0,
        metodo_pago_devolucion=venta.metodo_pago,
        usuario=request.user if request.user.is_authenticated else Usuario.objects.first(),
    )

    for i in range(len(producto_ids)):
        producto = get_object_or_404(Producto, pk=producto_ids[i])
        cantidad = int(cantidades[i])

        DetalleDevolucion.objects.create(
            devolucion=devolucion,
            producto=producto,
            cantidad=cantidad,
            fecha_vencimiento=timezone.now().date(),
            descripcion=f"Devolución de {producto.nombre}",
        )

    messages.success(request, f'Devolución #{devolucion.codigo_devolucion} registrada correctamente.')
    return redirect('comprobante_devolucion', pk=devolucion.pk)


@session_required
def comprobante_devolucion(request, pk):
    devolucion = get_object_or_404(Devolucion, pk=pk)
    return render(request, 'ventas/comprobante_devolucion.html', {'devolucion': devolucion})


# ════════════════════════════════════════
# EXPORTACIÓN EXCEL Y PDF DE VENTAS
# ════════════════════════════════════════

@session_required
def exportar_ventas_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.db.models import Q
    from django.http import HttpResponse

    fecha_inicio_str = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_str = request.GET.get('fecha_fin', '').strip()
    q = request.GET.get('q', '').strip()

    ventas_qs = Venta.objects.all().order_by('-fecha')

    if fecha_inicio_str and fecha_fin_str:
        try:
            f_ini = datetime.strptime(fecha_inicio_str, '%Y-%m-%d').date()
            f_fin = datetime.strptime(fecha_fin_str, '%Y-%m-%d').date()
            ventas_qs = ventas_qs.filter(fecha__date__range=[f_ini, f_fin])
        except ValueError:
            pass

    if q:
        ventas_qs = ventas_qs.filter(
            Q(pk__icontains=q) |
            Q(cliente__nombre__icontains=q) |
            Q(vendedor__first_name__icontains=q) |
            Q(vendedor__username__icontains=q) |
            Q(metodo_pago__icontains=q) |
            Q(detalles__producto__nombre__icontains=q)
        ).distinct()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"

    header_fill = PatternFill(start_color="0F3B6C", end_color="0F3B6C", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F3B6C")
    sub_font = Font(name="Calibri", size=10, italic=True, color="555555")
    border_thin = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    ws.merge_cells("A1:J1")
    ws["A1"] = "CYS Ltda — Reporte Consolidado de Ventas"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    gen_date = timezone.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A2:J2")
    ws["A2"] = f"Fecha de generación: {gen_date} | Transacciones: {ventas_qs.count()}"
    ws["A2"].font = sub_font

    ws.append([])

    headers = ["# Venta", "Fecha", "Cliente", "Vendedor", "Producto / Ítem", "Cantidad", "Precio Uni.", "Subtotal", "Método de Pago", "Total Venta"]
    ws.append(headers)
    header_row_idx = 4

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_acreditado = Decimal('0')

    for v in ventas_qs.select_related('vendedor').prefetch_related('detalles__producto'):
        detalles = list(v.detalles.all())
        cliente_nombre = v.cliente.nombre if v.cliente else "Consumidor Final"
        vendedor_nombre = str(v.vendedor.username if v.vendedor else "Sistema")
        fecha_str = v.fecha.strftime("%d/%m/%Y %H:%M")

        if detalles:
            for det in detalles:
                ws.append([
                    v.pk,
                    fecha_str,
                    cliente_nombre,
                    vendedor_nombre,
                    det.producto.nombre if det.producto else "Ítem",
                    det.cantidad,
                    float(det.precio_unitario),
                    float(det.subtotal),
                    (v.metodo_pago or 'Efectivo').capitalize(),
                    float(v.total_venta)
                ])
        else:
            ws.append([
                v.pk,
                fecha_str,
                cliente_nombre,
                vendedor_nombre,
                "—",
                0,
                0,
                0,
                (v.metodo_pago or 'Efectivo').capitalize(),
                float(v.total_venta)
            ])
        total_acreditado += v.total_venta

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.border = border_thin
            if cell.column in [7, 8, 10]:
                cell.number_format = '$#,##0'
            elif cell.column in [1, 6]:
                cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Reporte_Ventas_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    wb.save(response)
    return response


@session_required
def exportar_ventas_pdf(request):
    import io
    from django.db.models import Q
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    fecha_inicio_str = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_str = request.GET.get('fecha_fin', '').strip()
    q = request.GET.get('q', '').strip()

    ventas_qs = Venta.objects.all().order_by('-fecha')

    if fecha_inicio_str and fecha_fin_str:
        try:
            f_ini = datetime.strptime(fecha_inicio_str, '%Y-%m-%d').date()
            f_fin = datetime.strptime(fecha_fin_str, '%Y-%m-%d').date()
            ventas_qs = ventas_qs.filter(fecha__date__range=[f_ini, f_fin])
        except ValueError:
            pass

    if q:
        ventas_qs = ventas_qs.filter(
            Q(pk__icontains=q) |
            Q(cliente__nombre__icontains=q) |
            Q(vendedor__first_name__icontains=q) |
            Q(vendedor__username__icontains=q) |
            Q(metodo_pago__icontains=q) |
            Q(detalles__producto__nombre__icontains=q)
        ).distinct()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=18,
        textColor=colors.HexColor('#0F3B6C'), spaceAfter=4
    )
    sub_style = ParagraphStyle(
        'DocSub', parent=styles['Normal'],
        fontName='Helvetica', fontSize=10,
        textColor=colors.HexColor('#555555'), spaceAfter=15
    )
    cell_style = ParagraphStyle(
        'TableCell', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5,
        textColor=colors.HexColor('#222222')
    )
    cell_bold = ParagraphStyle(
        'TableCellBold', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=8.5,
        textColor=colors.HexColor('#0F3B6C')
    )
    cell_header = ParagraphStyle(
        'TableHeader', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=9,
        textColor=colors.white, alignment=1
    )

    elements = []
    elements.append(Paragraph("CYS Ltda — Reporte Consolidado de Ventas", title_style))
    gen_date = timezone.now().strftime("%d/%m/%Y %H:%M")
    elements.append(Paragraph(f"Fecha de generación: {gen_date} | Total transacciones: {ventas_qs.count()}", sub_style))

    data = [[
        Paragraph("#", cell_header),
        Paragraph("Fecha", cell_header),
        Paragraph("Cliente", cell_header),
        Paragraph("Producto / Ítem", cell_header),
        Paragraph("Cant.", cell_header),
        Paragraph("Precio Uni.", cell_header),
        Paragraph("Total Venta", cell_header),
        Paragraph("Método Pago", cell_header),
    ]]

    total_general = Decimal('0')

    for v in ventas_qs.select_related('vendedor').prefetch_related('detalles__producto'):
        detalles = list(v.detalles.all())
        cliente_nombre = v.cliente.nombre if v.cliente else "Consumidor Final"
        fecha_str = v.fecha.strftime("%d/%m/%Y %H:%M")
        total_general += v.total_venta

        if detalles:
            for det in detalles:
                data.append([
                    Paragraph(str(v.pk), cell_style),
                    Paragraph(fecha_str, cell_style),
                    Paragraph(cliente_nombre, cell_style),
                    Paragraph(det.producto.nombre if det.producto else "—", cell_style),
                    Paragraph(str(det.cantidad), cell_style),
                    Paragraph(f"${det.precio_unitario:,.0f}".replace(',', '.'), cell_style),
                    Paragraph(f"${v.total_venta:,.0f}".replace(',', '.'), cell_bold),
                    Paragraph((v.metodo_pago or 'Efectivo').capitalize(), cell_style),
                ])
        else:
            data.append([
                Paragraph(str(v.pk), cell_style),
                Paragraph(fecha_str, cell_style),
                Paragraph(cliente_nombre, cell_style),
                Paragraph("—", cell_style),
                Paragraph("0", cell_style),
                Paragraph("—", cell_style),
                Paragraph(f"${v.total_venta:,.0f}".replace(',', '.'), cell_bold),
                Paragraph((v.metodo_pago or 'Efectivo').capitalize(), cell_style),
            ])

    data.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph(f"<b>${total_general:,.0f}</b>".replace(',', '.'), cell_bold),
        Paragraph("", cell_style),
    ])

    col_widths = [35, 95, 130, 200, 45, 75, 85, 75]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F3B6C')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0, 1), (-1, -2), colors.HexColor('#FFFFFF')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E2F1FF')),
    ]))

    elements.append(table)
    doc.build(elements)

    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Reporte_Ventas_{timezone.now().strftime("%Y%m%d_%H%M")}.pdf"'
    response.write(pdf)
    return response
